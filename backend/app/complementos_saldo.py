from datetime import date
from io import BytesIO
from typing import Optional

import openpyxl
from openpyxl.styles import Font, PatternFill
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import Column, Date, DateTime, Float, Integer, String, func, or_
from sqlalchemy.orm import Session

from backend.app.database import Base, get_db
from backend.app.auth.dependencies import get_current_user
from backend.app.auth.permissions import require_perm_modulo
from backend.app.expedientes.models import Expediente

router = APIRouter(prefix="/complementos-saldo", tags=["Complementos de saldo"])

def _autorizar_complementos_ver(usuario=Depends(get_current_user)):
    for modulo in ("expedientes", "utilidades"):
        try:
            require_perm_modulo(usuario, modulo, "ver")
            return usuario
        except HTTPException as exc:
            if exc.status_code != 403:
                raise
    raise HTTPException(status_code=403, detail="No tienes permiso para consultar los complementos de saldo.")


def _autorizar_complementos_editar(usuario=Depends(get_current_user)):
    for modulo in ("expedientes", "utilidades"):
        try:
            require_perm_modulo(usuario, modulo, "editar")
            return usuario
        except HTTPException as exc:
            if exc.status_code != 403:
                raise
    raise HTTPException(status_code=403, detail="No tienes permiso para modificar los complementos de saldo.")


_autorizar_complementos_ver._molsan_permiso_guard = ("expedientes", "ver")
_autorizar_complementos_editar._molsan_permiso_guard = ("expedientes", "editar")



class ComplementoSaldo(Base):
    __tablename__ = "complementos_saldo"

    id = Column(Integer, primary_key=True, index=True)
    expediente_id = Column(Integer, nullable=True, index=True)
    numero_expediente = Column(String(100), nullable=False, index=True)
    fecha_solicitud = Column(Date, nullable=False, default=date.today, index=True)
    solicitante = Column(String(300), nullable=True)
    oficina = Column(String(200), nullable=True, index=True)
    importe = Column(Float, nullable=False, default=0)
    provisionado = Column(Float, nullable=False, default=0)
    saldo_disponible = Column(Float, nullable=False, default=0)
    motivo = Column(String(2000), nullable=True)
    estado = Column(String(20), nullable=False, default="Pendiente", index=True)
    observaciones = Column(String(2000), nullable=True)
    creado_por = Column(String(200), nullable=True)
    creado_en = Column(DateTime, nullable=False, server_default=func.now())


class ComplementoSaldoCrear(BaseModel):
    expediente_id: Optional[int] = None
    numero_expediente: str = Field(..., min_length=1, max_length=100)
    fecha_solicitud: date = Field(default_factory=date.today)
    solicitante: Optional[str] = None
    oficina: Optional[str] = None
    importe: float = Field(..., gt=0)
    provisionado: float = Field(default=0, ge=0)
    saldo_disponible: float = Field(default=0)
    motivo: str = Field(..., min_length=1, max_length=2000)
    estado: str = "Pendiente"
    observaciones: Optional[str] = None


def _serializar(item):
    return {
        "id": item.id,
        "expediente_id": item.expediente_id,
        "numero_expediente": item.numero_expediente,
        "fecha_solicitud": item.fecha_solicitud.isoformat() if item.fecha_solicitud else None,
        "solicitante": item.solicitante or "",
        "oficina": item.oficina or "",
        "importe": float(item.importe or 0),
        "provisionado": float(item.provisionado or 0),
        "saldo_disponible": float(item.saldo_disponible or 0),
        "motivo": item.motivo or item.observaciones or "",
        "estado": item.estado or "Pendiente",
        "observaciones": item.observaciones or item.motivo or "",
        "creado_por": item.creado_por or "",
        "creado_en": item.creado_en.isoformat() if item.creado_en else None,
    }


@router.get("", dependencies=[Depends(_autorizar_complementos_ver)])
def listar_complementos_saldo(
    q: Optional[str] = Query(None),
    estado: Optional[str] = Query(None),
    oficina: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    usuario=Depends(get_current_user),
):
    consulta = db.query(ComplementoSaldo)
    if q and q.strip():
        termino = f"%{q.strip()}%"
        consulta = consulta.filter(or_(
            ComplementoSaldo.numero_expediente.ilike(termino),
            ComplementoSaldo.solicitante.ilike(termino),
            ComplementoSaldo.oficina.ilike(termino),
            ComplementoSaldo.observaciones.ilike(termino),
        ))
    if estado and estado.lower() in ("pendiente", "compensado", "cerrado"):
        consulta = consulta.filter(func.lower(ComplementoSaldo.estado) == estado.lower())
    if oficina and oficina.strip():
        consulta = consulta.filter(ComplementoSaldo.oficina.ilike(f"%{oficina.strip()}%"))
    items = consulta.order_by(ComplementoSaldo.fecha_solicitud.desc(), ComplementoSaldo.id.desc()).all()
    return [_serializar(item) for item in items]


@router.post("", status_code=201, dependencies=[Depends(_autorizar_complementos_editar)])
def crear_complemento_saldo(
    datos: ComplementoSaldoCrear,
    db: Session = Depends(get_db),
    usuario=Depends(get_current_user),
):
    numero = datos.numero_expediente.strip()
    expediente = None
    if datos.expediente_id is not None:
        expediente = db.query(Expediente).filter(Expediente.id == datos.expediente_id).first()
    if expediente is None:
        expediente = db.query(Expediente).filter(Expediente.id_expediente == numero).first()
    if expediente is not None:
        numero = expediente.id_expediente or numero
    estado = "Cerrado" if str(datos.estado).lower() == "cerrado" else "Pendiente"
    nombre_usuario = (
        getattr(usuario, "nombre_completo", None)
        or getattr(usuario, "nombre", None)
        or getattr(usuario, "email", None)
        or getattr(usuario, "username", None)
    )
    item = ComplementoSaldo(
        expediente_id=expediente.id if expediente else datos.expediente_id,
        numero_expediente=numero,
        fecha_solicitud=datos.fecha_solicitud,
        solicitante=(datos.solicitante or (getattr(expediente, "nombre_solicitante", None) if expediente else None) or "").strip(),
        oficina=(datos.oficina or (getattr(expediente, "oficina", None) if expediente else None) or "").strip(),
        importe=datos.importe,
        provisionado=datos.provisionado,
        saldo_disponible=datos.saldo_disponible,
        motivo=(datos.motivo or "").strip(),
        estado=estado,
        observaciones=datos.observaciones or datos.motivo,
        creado_por=str(nombre_usuario or ""),
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    return _serializar(item)


@router.patch("/{complemento_id}/estado", dependencies=[Depends(_autorizar_complementos_editar)])
def cambiar_estado_complemento_saldo(
    complemento_id: int,
    estado: str,
    db: Session = Depends(get_db),
    usuario=Depends(get_current_user),
):
    if estado.lower() not in ("pendiente", "compensado", "cerrado"):
        raise HTTPException(status_code=400, detail="El estado debe ser Pendiente o Compensado.")
    item = db.query(ComplementoSaldo).filter(ComplementoSaldo.id == complemento_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="No se ha encontrado la solicitud de complemento.")
    item.estado = "Compensado" if estado.lower() in ("compensado", "cerrado") else "Pendiente"
    db.commit()
    db.refresh(item)
    return _serializar(item)


@router.get("/exportar/excel", dependencies=[Depends(_autorizar_complementos_ver)])
def exportar_complementos_saldo(
    q: Optional[str] = Query(None),
    estado: Optional[str] = Query(None),
    oficina: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    usuario=Depends(get_current_user),
):
    items = listar_complementos_saldo(q=q, estado=estado, oficina=oficina, db=db, usuario=usuario)
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Complementos de saldo"
    hoja.append(["Expediente", "Provisionado (€)", "Saldo disponible (€)", "Fecha de petición", "Saldo a solicitar (€)", "Motivo", "Acción", "Estado", "Solicitado por"])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1D4ED8")
    for item in items:
        hoja.append([item["numero_expediente"], item["provisionado"], item["saldo_disponible"], item["fecha_solicitud"], item["importe"], item["motivo"], "Regularizado" if str(item["estado"]).lower() in ("compensado", "cerrado") else "Compensar", "Compensado" if str(item["estado"]).lower() in ("compensado", "cerrado") else "Pendiente", item["creado_por"]])
    for fila in hoja.iter_rows(min_row=2):
        for indice_importe in (2, 3, 5):
            fila[indice_importe].number_format = '#,##0.00 [$€-es-ES]'
    for columna in hoja.columns:
        letra = columna[0].column_letter
        hoja.column_dimensions[letra].width = min(max(max(len(str(c.value or "")) for c in columna) + 2, 14), 38)
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    nombre = f"complementos_saldo_{date.today().isoformat()}.xlsx"
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{nombre}"'})
