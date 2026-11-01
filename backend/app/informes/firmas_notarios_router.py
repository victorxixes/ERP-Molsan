from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Optional
import unicodedata

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy import or_, func
from sqlalchemy.orm import Session

from backend.app.database import get_db
from backend.app.auth.dependencies import get_current_user
from backend.app.auth.permissions import require_perm_modulo
from backend.app.expedientes.models import Expediente

def _autorizar_informes(usuario=Depends(get_current_user)):
    require_perm_modulo(usuario, "informes", "ver")
    return usuario


def _autorizar_edicion_expedientes(usuario=Depends(get_current_user)):
    require_perm_modulo(usuario, "expedientes", "editar")
    return usuario


router = APIRouter(
    prefix="/informes",
    tags=["Informes"],
    dependencies=[Depends(get_current_user), Depends(_autorizar_informes)],
)


def _normalizar(texto) -> str:
    valor = unicodedata.normalize("NFKD", str(texto or ""))
    return "".join(c for c in valor if not unicodedata.combining(c)).strip().lower()


def _clasificar_modalidad(tipo_firma) -> str:
    valor = _normalizar(tipo_firma)
    compacto = valor.replace(" ", "").replace("-", "_")
    if not valor:
        return "SIN_CLASIFICAR"
    if (
        compacto in {"vc", "video", "videoconferencia", "videoconference"}
        or "videoconferencia" in compacto
        or "videoconference" in compacto
        or "video conferencia" in valor
    ):
        return "VC"
    if "presencial" in compacto or compacto in {"p", "fisica"}:
        return "PRESENCIAL"
    return "SIN_CLASIFICAR"


def _fecha_envio(expediente):
    return getattr(expediente, "fecha_envio_notario", None) or getattr(expediente, "fecha_envio", None)


def _dias_envio_firma(expediente) -> Optional[int]:
    envio = _fecha_envio(expediente)
    firma = getattr(expediente, "fecha_firma", None)
    if not envio or not firma:
        return None
    dias = (firma - envio).days
    return dias if dias >= 0 else None


def _firmas_query(
    db: Session,
    fecha_desde: Optional[date] = None,
    fecha_hasta: Optional[date] = None,
    modalidad: Optional[str] = None,
    tipo_operacion: Optional[str] = None,
    notario: Optional[str] = None,
    apoderado: Optional[str] = None,
    q: Optional[str] = None,
):
    if fecha_desde and fecha_hasta and fecha_desde > fecha_hasta:
        raise HTTPException(status_code=400, detail="La fecha desde no puede ser posterior a la fecha hasta.")

    modalidad_normalizada = _normalizar(modalidad).replace(" ", "_")
    modalidades_validas = {
        "": None,
        "vc": "VC",
        "videoconferencia": "VC",
        "presencial": "PRESENCIAL",
        "sin_clasificar": "SIN_CLASIFICAR",
        "sin clasificar": "SIN_CLASIFICAR",
    }
    if modalidad_normalizada not in modalidades_validas:
        raise HTTPException(status_code=400, detail="Modalidad no válida. Usa VC, PRESENCIAL o SIN_CLASIFICAR.")
    modalidad_filtro = modalidades_validas[modalidad_normalizada]

    query = db.query(Expediente).filter(Expediente.fecha_firma.isnot(None))

    if fecha_desde:
        query = query.filter(Expediente.fecha_firma >= fecha_desde)
    if fecha_hasta:
        query = query.filter(Expediente.fecha_firma <= fecha_hasta)
    if tipo_operacion and tipo_operacion.strip():
        patron = f"%{tipo_operacion.strip()}%"
        query = query.filter(Expediente.tipo_operacion.ilike(patron))
    if notario and notario.strip():
        patron = f"%{notario.strip()}%"
        query = query.filter(or_(
            Expediente.nombre_notario.ilike(patron),
            Expediente.notario.ilike(patron),
            Expediente.nif_notario.ilike(patron),
        ))
    if apoderado and apoderado.strip():
        query = query.filter(Expediente.apoderado.ilike(f"%{apoderado.strip()}%"))
    if q and q.strip():
        patron = f"%{q.strip()}%"
        query = query.filter(or_(
            Expediente.id_expediente.ilike(patron),
            Expediente.nombre_titular.ilike(patron),
            Expediente.nif_titular.ilike(patron),
            Expediente.nombre_solicitante.ilike(patron),
            Expediente.nif_solicitante.ilike(patron),
            Expediente.protocolo.ilike(patron),
            Expediente.oficina.ilike(patron),
            Expediente.nombre_notario.ilike(patron),
            Expediente.notario.ilike(patron),
        ))

    expedientes = query.order_by(Expediente.fecha_firma.desc(), Expediente.id.desc()).all()
    items = []
    for exp in expedientes:
        modo = _clasificar_modalidad(getattr(exp, "tipo_firma", None))
        if modalidad_filtro and modo != modalidad_filtro:
            continue
        envio = _fecha_envio(exp)
        dias = _dias_envio_firma(exp)
        items.append({
            "id": exp.id,
            "id_expediente": getattr(exp, "id_expediente", None),
            "fecha_firma": exp.fecha_firma.isoformat() if exp.fecha_firma else None,
            "fecha_envio": envio.isoformat() if envio else None,
            "dias_envio_firma": dias,
            "tipo_firma": getattr(exp, "tipo_firma", None),
            "modalidad": modo,
            "tipo_operacion": getattr(exp, "tipo_operacion", None),
            "nombre_notario": getattr(exp, "nombre_notario", None) or getattr(exp, "notario", None),
            "nif_notario": getattr(exp, "nif_notario", None),
            "apoderado": getattr(exp, "apoderado", None),
            "protocolo": getattr(exp, "protocolo", None),
            "nombre_titular": getattr(exp, "nombre_titular", None),
            "nif_titular": getattr(exp, "nif_titular", None),
            "nombre_solicitante": getattr(exp, "nombre_solicitante", None),
            "nif_solicitante": getattr(exp, "nif_solicitante", None),
            "oficina": getattr(exp, "oficina", None),
            "importe": getattr(exp, "importe", None),
            "capital": getattr(exp, "capital", None),
            "estado_expediente": getattr(exp, "estado_expediente", None),
            "actividad_actual": getattr(exp, "actividad_actual", None),
            "observaciones": getattr(exp, "observaciones", None),
        })
    return items


def _resumen_firmas(items: list[dict]) -> dict:
    total = len(items)
    vc = sum(1 for x in items if x["modalidad"] == "VC")
    presencial = sum(1 for x in items if x["modalidad"] == "PRESENCIAL")
    sin_clasificar = sum(1 for x in items if x["modalidad"] == "SIN_CLASIFICAR")
    dias = [x["dias_envio_firma"] for x in items if x["dias_envio_firma"] is not None]

    por_operacion = {}
    for item in items:
        nombre = (item.get("tipo_operacion") or "Sin tipo de operación").strip() or "Sin tipo de operación"
        grupo = por_operacion.setdefault(nombre, {"tipo_operacion": nombre, "total_firmas": 0, "vc": 0, "presencial": 0, "sin_clasificar": 0, "_dias": []})
        grupo["total_firmas"] += 1
        if item["modalidad"] == "VC":
            grupo["vc"] += 1
        elif item["modalidad"] == "PRESENCIAL":
            grupo["presencial"] += 1
        else:
            grupo["sin_clasificar"] += 1
        if item["dias_envio_firma"] is not None:
            grupo["_dias"].append(item["dias_envio_firma"])

    operaciones = []
    for grupo in por_operacion.values():
        muestras = grupo.pop("_dias")
        grupo["media_dias"] = round(sum(muestras) / len(muestras), 1) if muestras else None
        operaciones.append(grupo)
    operaciones.sort(key=lambda x: (-x["total_firmas"], x["tipo_operacion"].lower()))

    return {
        "total_firmas": total,
        "videoconferencia": vc,
        "presencial": presencial,
        "sin_clasificar": sin_clasificar,
        "firmas_con_tiempo_calculable": len(dias),
        "media_dias_envio_firma": round(sum(dias) / len(dias), 1) if dias else None,
        "por_tipo_operacion": operaciones,
    }


def _respuesta_excel(nombre: str, columnas: list[tuple[str, str]], filas: list[dict], filename: str):
    wb = Workbook()
    ws = wb.active
    ws.title = nombre[:31]
    ws.freeze_panes = "A2"
    for indice, (cabecera, _) in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=indice, value=cabecera)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1D4ED8")
        celda.alignment = Alignment(vertical="center")
    for numero_fila, fila in enumerate(filas, start=2):
        for indice, (_, campo) in enumerate(columnas, start=1):
            valor = fila.get(campo)
            ws.cell(row=numero_fila, column=indice, value=valor)
    if filas:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{len(filas) + 1}"
    for indice, (cabecera, campo) in enumerate(columnas, start=1):
        maximo = len(cabecera)
        for fila in filas[:500]:
            valor = fila.get(campo)
            if valor is not None:
                maximo = max(maximo, len(str(valor)))
        ws.column_dimensions[get_column_letter(indice)].width = min(max(maximo + 2, 12), 42)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


COLUMNAS_FIRMAS_EXCEL = [
    ("Nº expediente", "id_expediente"), ("Fecha firma", "fecha_firma"),
    ("Fecha envío", "fecha_envio"), ("Días envío-firma", "dias_envio_firma"),
    ("Modalidad", "modalidad"), ("Tipo firma original", "tipo_firma"),
    ("Tipo operación", "tipo_operacion"), ("Notario", "nombre_notario"),
    ("NIF notario", "nif_notario"), ("Apoderado", "apoderado"),
    ("Titular", "nombre_titular"), ("NIF titular", "nif_titular"),
    ("Solicitante", "nombre_solicitante"), ("NIF solicitante", "nif_solicitante"),
    ("Protocolo", "protocolo"), ("Oficina", "oficina"),
    ("Importe", "importe"), ("Capital", "capital"),
    ("Estado expediente", "estado_expediente"), ("Actividad actual", "actividad_actual"),
    ("Observaciones", "observaciones"),
]

def _params_firmas(
    fecha_desde: Optional[date], fecha_hasta: Optional[date], modalidad: Optional[str],
    tipo_operacion: Optional[str], notario: Optional[str], apoderado: Optional[str], q: Optional[str],
):
    return dict(fecha_desde=fecha_desde, fecha_hasta=fecha_hasta, modalidad=modalidad,
                tipo_operacion=tipo_operacion, notario=notario, apoderado=apoderado, q=q)


@router.get("/firmas")
def informe_firmas(
    fecha_desde: Optional[date] = Query(None),
    fecha_hasta: Optional[date] = Query(None),
    modalidad: Optional[str] = Query(None),
    tipo_operacion: Optional[str] = Query(None),
    notario: Optional[str] = Query(None),
    apoderado: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    """Informe de firmas reales. Solo cuenta expedientes con fecha_firma; no usa citas de Agenda."""
    items = _firmas_query(db, **_params_firmas(fecha_desde, fecha_hasta, modalidad, tipo_operacion, notario, apoderado, q))
    resumen = _resumen_firmas(items)
    return {
        "periodo": {"fecha_desde": fecha_desde.isoformat() if fecha_desde else None,
                    "fecha_hasta": fecha_hasta.isoformat() if fecha_hasta else None},
        "resumen": resumen,
        "por_modalidad": [
            {"modalidad": "VC", "total": resumen["videoconferencia"]},
            {"modalidad": "PRESENCIAL", "total": resumen["presencial"]},
            {"modalidad": "SIN_CLASIFICAR", "total": resumen["sin_clasificar"]},
        ],
        "por_tipo_operacion": resumen["por_tipo_operacion"],
        "items": items,
    }


class ActualizarDatosFirmaRequest(BaseModel):
    """Corrección verificada de los datos de una firma histórica."""
    modalidad: Optional[str] = None
    fecha_envio: Optional[date] = None


@router.put("/firmas/{id_expediente}/datos", dependencies=[Depends(_autorizar_edicion_expedientes)])
def actualizar_datos_firma(
    id_expediente: str,
    datos: ActualizarDatosFirmaRequest,
    db: Session = Depends(get_db),
):
    """Permite completar modalidad real y fecha de envío de un expediente firmado.

    No crea firmas ni consulta Agenda. Solo modifica un expediente que ya tiene
    fecha_firma. Los datos deben introducirse a partir de documentación verificada.
    """
    campos_enviados = getattr(datos, "__fields_set__", set())
    if not campos_enviados.intersection({"modalidad", "fecha_envio"}):
        raise HTTPException(status_code=400, detail="Indica la modalidad y/o la fecha de envío que quieres corregir.")

    expediente = (
        db.query(Expediente)
        .filter(Expediente.id_expediente == str(id_expediente))
        .first()
    )
    if expediente is None:
        raise HTTPException(status_code=404, detail="Expediente no encontrado.")
    if not getattr(expediente, "fecha_firma", None):
        raise HTTPException(status_code=409, detail="Solo se pueden corregir datos de un expediente con fecha de firma registrada.")

    # Valida las fechas antes de mutar el objeto, para que un error de
    # validación no deje cambios pendientes en la sesión.
    if "fecha_envio" in campos_enviados and datos.fecha_envio:
        if expediente.fecha_firma and datos.fecha_envio > expediente.fecha_firma:
            raise HTTPException(status_code=400, detail="La fecha de envío no puede ser posterior a la fecha de firma.")

    modalidad_guardada = getattr(expediente, "tipo_firma", None)
    fecha_envio_guardada = _fecha_envio(expediente)

    if "modalidad" in campos_enviados:
        valor = _normalizar(datos.modalidad).replace(" ", "_")
        if valor in {"vc", "video", "videoconferencia", "videoconference"}:
            modalidad_guardada = "Videoconferencia"
        elif valor in {"p", "presencial", "fisica", "physical"}:
            modalidad_guardada = "Presencial"
        elif valor in {"", "sin_clasificar", "no_informada", "sin_dato"}:
            modalidad_guardada = None
        else:
            raise HTTPException(status_code=400, detail="La modalidad debe ser VC, PRESENCIAL o SIN_CLASIFICAR.")
        expediente.tipo_firma = modalidad_guardada

    if "fecha_envio" in campos_enviados:
        fecha_envio_guardada = datos.fecha_envio
        # Ambos campos existen en el modelo: el primero es el dato del envío
        # al notario; el segundo se utiliza en las actividades del expediente.
        expediente.fecha_envio_notario = fecha_envio_guardada
        expediente.fecha_envio = fecha_envio_guardada

    try:
        db.add(expediente)
        db.commit()
        db.refresh(expediente)
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"No se pudieron guardar los datos de la firma: {str(exc)[:300]}")

    envio = _fecha_envio(expediente)
    dias = _dias_envio_firma(expediente)
    modalidad_resultante = _clasificar_modalidad(getattr(expediente, "tipo_firma", None))
    return {
        "ok": True,
        "id_expediente": expediente.id_expediente,
        "fecha_firma": expediente.fecha_firma.isoformat() if expediente.fecha_firma else None,
        "fecha_envio": envio.isoformat() if envio else None,
        "dias_envio_firma": dias,
        "tipo_firma": getattr(expediente, "tipo_firma", None),
        "modalidad": modalidad_resultante,
    }


@router.get("/firmas/exportar-excel")
def exportar_informe_firmas_excel(
    fecha_desde: Optional[date] = Query(None),
    fecha_hasta: Optional[date] = Query(None),
    modalidad: Optional[str] = Query(None),
    tipo_operacion: Optional[str] = Query(None),
    notario: Optional[str] = Query(None),
    apoderado: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    items = _firmas_query(db, **_params_firmas(fecha_desde, fecha_hasta, modalidad, tipo_operacion, notario, apoderado, q))
    return _respuesta_excel("Informe firmas", COLUMNAS_FIRMAS_EXCEL, items, "informe_firmas.xlsx")




# ============================================================
# INFORME DE NOTARIOS — FIRMAS REALES
# ============================================================

def _clave_notario_item(item: dict) -> str:
    nif = str(item.get("nif_notario") or "").strip().upper()
    if nif:
        return f"NIF:{nif}"
    nombre = _normalizar(item.get("nombre_notario") or item.get("notario") or "")
    if nombre:
        return f"NOMBRE:{nombre}"
    return "SIN_NOTARIO_IDENTIFICADO"


def _resumen_por_notario(items: list[dict]) -> list[dict]:
    grupos: dict[str, dict] = {}
    for item in items:
        clave = _clave_notario_item(item)
        item["clave_notario"] = clave
        nif = str(item.get("nif_notario") or "").strip().upper() or None
        nombre = str(item.get("nombre_notario") or item.get("notario") or "").strip()
        if not nombre:
            nombre = f"Notario sin nombre ({nif})" if nif else "Sin notario identificado"

        grupo = grupos.get(clave)
        if grupo is None:
            grupo = {
                "clave_notario": clave,
                "nombre_notario": nombre,
                "nif_notario": nif,
                "total_firmas": 0,
                "vc": 0,
                "presencial": 0,
                "sin_clasificar": 0,
                "firmas_con_tiempo_calculable": 0,
                "_dias": [],
                "_fechas": [],
                "_operaciones": set(),
                "_apoderados": set(),
                "con_protocolo": 0,
            }
            grupos[clave] = grupo
        elif not grupo.get("nif_notario") and nif:
            grupo["nif_notario"] = nif
        if nombre and (not grupo["nombre_notario"] or grupo["nombre_notario"] == "Sin notario identificado"):
            grupo["nombre_notario"] = nombre

        grupo["total_firmas"] += 1
        modalidad = item.get("modalidad")
        if modalidad == "VC":
            grupo["vc"] += 1
        elif modalidad == "PRESENCIAL":
            grupo["presencial"] += 1
        else:
            grupo["sin_clasificar"] += 1

        dias = item.get("dias_envio_firma")
        if dias is not None:
            grupo["_dias"].append(int(dias))
            grupo["firmas_con_tiempo_calculable"] += 1

        fecha_firma = item.get("fecha_firma")
        if fecha_firma:
            grupo["_fechas"].append(str(fecha_firma)[:10])
        operacion = str(item.get("tipo_operacion") or "").strip()
        if operacion:
            grupo["_operaciones"].add(operacion)
        apoderado = str(item.get("apoderado") or "").strip()
        if apoderado:
            grupo["_apoderados"].add(apoderado)
        if str(item.get("protocolo") or "").strip():
            grupo["con_protocolo"] += 1

    resultado = []
    for grupo in grupos.values():
        dias = grupo.pop("_dias")
        fechas = grupo.pop("_fechas")
        operaciones = grupo.pop("_operaciones")
        apoderados = grupo.pop("_apoderados")
        grupo["media_dias_envio_firma"] = round(sum(dias) / len(dias), 1) if dias else None
        grupo["ultima_firma"] = max(fechas) if fechas else None
        grupo["tipos_operacion"] = sorted(operaciones, key=str.lower)
        grupo["numero_tipos_operacion"] = len(operaciones)
        grupo["apoderados"] = sorted(apoderados, key=str.lower)
        grupo["numero_apoderados"] = len(apoderados)
        resultado.append(grupo)

    resultado.sort(key=lambda fila: (-fila["total_firmas"], _normalizar(fila["nombre_notario"])))
    return resultado


def _resumen_notarios(items: list[dict], por_notario: list[dict]) -> dict:
    resumen_firmas = _resumen_firmas(items)
    fechas = [str(item.get("fecha_firma"))[:10] for item in items if item.get("fecha_firma")]
    return {
        "total_firmas": int(resumen_firmas["total_firmas"]),
        "total_notarios": sum(1 for fila in por_notario if fila["clave_notario"] != "SIN_NOTARIO_IDENTIFICADO"),
        "videoconferencia": int(resumen_firmas["videoconferencia"]),
        "presencial": int(resumen_firmas["presencial"]),
        "sin_clasificar": int(resumen_firmas["sin_clasificar"]),
        "firmas_con_tiempo_calculable": int(resumen_firmas["firmas_con_tiempo_calculable"]),
        "media_dias_envio_firma": resumen_firmas["media_dias_envio_firma"],
        "con_protocolo": sum(1 for item in items if str(item.get("protocolo") or "").strip()),
        "ultima_firma": max(fechas) if fechas else None,
        "sin_notario_identificado": sum(1 for fila in por_notario if fila["clave_notario"] == "SIN_NOTARIO_IDENTIFICADO"),
    }


COLUMNAS_NOTARIOS_EXCEL = [
    ("Notario", "nombre_notario"),
    ("NIF notario", "nif_notario"),
    ("Expedientes firmados", "total_firmas"),
    ("Videoconferencia", "vc"),
    ("Presencial", "presencial"),
    ("Sin clasificar", "sin_clasificar"),
    ("Con protocolo", "con_protocolo"),
    ("Firmas con tiempo calculable", "firmas_con_tiempo_calculable"),
    ("Media días envío-firma", "media_dias_envio_firma"),
    ("Última firma", "ultima_firma"),
    ("Tipos de operación", "tipos_operacion_texto"),
    ("Número de apoderados", "numero_apoderados"),
]

COLUMNAS_NOTARIOS_DETALLE_EXCEL = [
    ("Nº expediente", "id_expediente"), ("Fecha firma", "fecha_firma"),
    ("Fecha envío", "fecha_envio"), ("Días envío-firma", "dias_envio_firma"),
    ("Notario", "nombre_notario"), ("NIF notario", "nif_notario"),
    ("Modalidad", "modalidad"), ("Tipo operación", "tipo_operacion"),
    ("Apoderado", "apoderado"), ("Protocolo", "protocolo"),
    ("Titular", "nombre_titular"), ("NIF titular", "nif_titular"),
    ("Oficina", "oficina"), ("Importe", "importe"), ("Capital", "capital"),
]


def _params_notarios_excel(por_notario: list[dict]) -> list[dict]:
    filas = []
    for fila in por_notario:
        copia = dict(fila)
        copia["tipos_operacion_texto"] = ", ".join(fila.get("tipos_operacion") or [])
        filas.append(copia)
    return filas


@router.get("/notarios")
def informe_notarios(
    fecha_desde: Optional[date] = Query(None),
    fecha_hasta: Optional[date] = Query(None),
    modalidad: Optional[str] = Query(None),
    tipo_operacion: Optional[str] = Query(None),
    notario: Optional[str] = Query(None),
    apoderado: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    """Resume por notario únicamente expedientes con fecha_firma real."""
    items = _firmas_query(
        db,
        **_params_firmas(fecha_desde, fecha_hasta, modalidad, tipo_operacion, notario, apoderado, q),
    )
    por_notario = _resumen_por_notario(items)
    resumen = _resumen_notarios(items, por_notario)
    return {
        "periodo": {
            "fecha_desde": fecha_desde.isoformat() if fecha_desde else None,
            "fecha_hasta": fecha_hasta.isoformat() if fecha_hasta else None,
        },
        "resumen": resumen,
        "por_notario": por_notario,
        "items": items,
    }


@router.get("/notarios/exportar-excel")
def exportar_informe_notarios_excel(
    fecha_desde: Optional[date] = Query(None),
    fecha_hasta: Optional[date] = Query(None),
    modalidad: Optional[str] = Query(None),
    tipo_operacion: Optional[str] = Query(None),
    notario: Optional[str] = Query(None),
    apoderado: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    items = _firmas_query(
        db,
        **_params_firmas(fecha_desde, fecha_hasta, modalidad, tipo_operacion, notario, apoderado, q),
    )
    grupos = _resumen_por_notario(items)
    filas = _params_notarios_excel(grupos)
    return _respuesta_excel("Informe notarios", COLUMNAS_NOTARIOS_EXCEL, filas, "informe_notarios.xlsx")


@router.get("/notarios/exportar-detalle-excel")
def exportar_detalle_informe_notarios_excel(
    fecha_desde: Optional[date] = Query(None),
    fecha_hasta: Optional[date] = Query(None),
    modalidad: Optional[str] = Query(None),
    tipo_operacion: Optional[str] = Query(None),
    notario: Optional[str] = Query(None),
    apoderado: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    items = _firmas_query(
        db,
        **_params_firmas(fecha_desde, fecha_hasta, modalidad, tipo_operacion, notario, apoderado, q),
    )
    return _respuesta_excel("Detalle firmas", COLUMNAS_NOTARIOS_DETALLE_EXCEL, items, "detalle_firmas_notarios.xlsx")
