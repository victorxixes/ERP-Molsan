from __future__ import annotations

import json

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, or_
from sqlalchemy.orm import Session
from sqlalchemy import text
from backend.app.database import engine

from backend.app.auth.dependencies import get_current_user
from backend.app.database import get_db
from backend.app.expedientes.models import Expediente
from backend.app.expedientes.gastos.models import ExpedienteGasto

router = APIRouter(prefix="/informes", tags=["Informes"], dependencies=[Depends(get_current_user)])


@router.on_event("startup")
def migrar_campos_absis_facturacion():
    """Añade los campos opcionales del informe sin modificar datos existentes."""
    columnas = {
        "estado_absis": "VARCHAR(100)",
        "estado_justificante": "VARCHAR(100)",
        "fecha_adjuntar_factura": "DATE",
        "imp_gasto": "NUMERIC(12, 2)",
        "imp_base_impuesto": "NUMERIC(12, 2)",
        "imp_base_retencion": "NUMERIC(12, 2)",
        "imp_impuesto": "NUMERIC(12, 2)",
        "fecha_estado": "DATE",
        "fecha_peticion": "DATE",
        "fecha_pago": "DATE",
        "id_gasto_absis": "VARCHAR(100)",
        "id_expediente_absis": "VARCHAR(100)",
        "tipo_operacion_absis": "VARCHAR(300)",
        "concepto_gasto_absis": "VARCHAR(300)",
        "datos_absis_json": "TEXT",
        "numero_factura_absis": "VARCHAR(150)",
        "fecha_factura_absis": "DATE",
        "tipo_concepto_absis": "VARCHAR(100)",
        "origen_registro": "VARCHAR(30)",
        "estado_consolidacion": "VARCHAR(50)",
        "fecha_ultima_sincronizacion": "DATE",
    }
    with engine.begin() as conexion:
        for nombre, tipo_sql in columnas.items():
            conexion.execute(text(
                f"ALTER TABLE public.expediente_gastos ADD COLUMN IF NOT EXISTS {nombre} {tipo_sql}"
            ))
        conexion.execute(text("CREATE INDEX IF NOT EXISTS ix_expediente_gastos_id_gasto_absis ON public.expediente_gastos (id_gasto_absis)"))
        conexion.execute(text("CREATE INDEX IF NOT EXISTS ix_expediente_gastos_id_expediente_absis ON public.expediente_gastos (id_expediente_absis)"))

def _normalizar_texto(valor):
    import unicodedata
    texto = str(valor or "").strip().casefold()
    return "".join(c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c))

def _concepto_honorarios(gasto):
    # Compatibilidad con datos anteriores: el concepto puede estar en tipo o descripcion.
    valores = [getattr(gasto, "tipo", None), getattr(gasto, "descripcion", None)]
    return any("honorarios gestoria" in _normalizar_texto(v) for v in valores if v)

def _tipo_cancelacion(exp, gasto=None):
    valor_origen = getattr(gasto, "tipo_operacion_absis", None) if gasto is not None else None
    valor = _normalizar_texto(valor_origen or getattr(exp, "tipo_operacion", None))
    # ABSIS puede informar simplemente "CANCELACION"; se agrupa con las
    # cancelaciones con provisión, salvo que indique expresamente lo contrario.
    if "cancelacion" in valor and "sin provision" in valor:
        return "sin_provision"
    if "cancelacion" in valor:
        return "con_provision"
    return "otros"

def _situacion_facturacion(gasto):
    """La situación económica se decide por ESTADO de ABSIS; FPAGO solo aporta la fecha."""
    estado = _normalizar_texto(getattr(gasto, "estado_absis", None))
    justificante = _normalizar_texto(getattr(gasto, "estado_justificante", None))
    conjunto = f"{estado} {justificante}"

    # Incidencias prevalecen para no contabilizar como cobro confirmado un registro problemático.
    if any(p in conjunto for p in ("incidencia", "error", "rechaz", "devuelt", "problema", "anulad")):
        return "incidencia"

    # No deducir el cobro de la existencia de FPAGO: el estado oficial manda.
    if any(p in estado for p in ("pagad", "abonad", "cobrad")):
        return "cobrado"

    # Si ABSIS aún no confirma el pago, clasificar según el estado real.
    if any(p in estado for p in ("pendient", "solicit", "tramite", "en curso", "emitid", "generad")):
        return "pendiente_cobro"

    # La fecha de factura sirve para distinguir pendiente de factura, nunca para declarar un cobro.
    if not estado and not justificante:
        return "pendiente_factura" if not getattr(gasto, "fecha_adjuntar_factura", None) else "pendiente_cobro"

    return "otros"

def _fecha_original_absis(gasto, clave):
    """Recupera una fecha del JSON original cuando el campo normalizado está vacío."""
    bruto = getattr(gasto, "datos_absis_json", None)
    if not bruto:
        return None
    try:
        original = json.loads(bruto)
    except (TypeError, ValueError):
        return None
    valor = original.get(clave)
    if valor is None or str(valor).strip() == "":
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor).strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(texto[:10], formato).date()
        except ValueError:
            continue
    return None


def _fecha_iso(valor):
    return valor.isoformat() if valor else None

def _decimal(valor):
    if valor is None:
        return None
    try:
        return float(Decimal(str(valor)))
    except Exception:
        return None

def _consultar(db, fecha_desde=None, fecha_hasta=None, tipo_operacion=None, estado=None, estado_justificante=None, situacion=None, q=None, modalidad="todas"):
    if fecha_desde and fecha_hasta and fecha_desde > fecha_hasta:
        raise HTTPException(status_code=400, detail="La fecha desde no puede ser posterior a la fecha hasta.")
    query = db.query(ExpedienteGasto, Expediente).outerjoin(
        Expediente, Expediente.id == ExpedienteGasto.expediente_id
    )
    # Para registros Post Firma, el concepto original de ABSIS es la fuente de verdad.
    # Solo los gastos ERP antiguos sin IDGASTO ABSIS pueden usar tipo/descripcion como compatibilidad.
    query = query.filter(or_(
        ExpedienteGasto.concepto_gasto_absis.ilike("%honorarios gestoría%"),
        ExpedienteGasto.concepto_gasto_absis.ilike("%honorarios gestoria%"),
        and_(
            ExpedienteGasto.id_gasto_absis.is_(None),
            or_(
                ExpedienteGasto.tipo.ilike("%honorarios gestoría%"),
                ExpedienteGasto.tipo.ilike("%honorarios gestoria%"),
                ExpedienteGasto.descripcion.ilike("%honorarios gestoría%"),
                ExpedienteGasto.descripcion.ilike("%honorarios gestoria%"),
            ),
        ),
    ))
    if fecha_desde:
        query = query.filter(ExpedienteGasto.fecha_adjuntar_factura >= fecha_desde)
    if fecha_hasta:
        query = query.filter(ExpedienteGasto.fecha_adjuntar_factura <= fecha_hasta)
    if tipo_operacion:
        query = query.filter(or_(Expediente.tipo_operacion.ilike(f"%{tipo_operacion.strip()}%"), ExpedienteGasto.tipo_operacion_absis.ilike(f"%{tipo_operacion.strip()}%")))
    if estado:
        query = query.filter(ExpedienteGasto.estado_absis.ilike(f"%{estado.strip()}%"))
    if estado_justificante:
        query = query.filter(ExpedienteGasto.estado_justificante.ilike(f"%{estado_justificante.strip()}%"))
    if q and q.strip():
        patron = f"%{q.strip()}%"
        query = query.filter(or_(
            Expediente.id_expediente.ilike(patron),
            ExpedienteGasto.id_expediente_absis.ilike(patron),
            Expediente.tipo_operacion.ilike(patron),
            ExpedienteGasto.tipo_operacion_absis.ilike(patron),
            ExpedienteGasto.concepto_gasto_absis.ilike(patron),
            ExpedienteGasto.tipo.ilike(patron),
            ExpedienteGasto.descripcion.ilike(patron),
            ExpedienteGasto.numero_factura.ilike(patron),
        ))
    filas = []
    for gasto, exp in query.order_by(ExpedienteGasto.fecha_adjuntar_factura.desc().nullslast(), Expediente.id_expediente.asc()).all():
        tipo = _tipo_cancelacion(exp, gasto)
        if modalidad in {"con_provision", "sin_provision", "otros"} and tipo != modalidad:
            continue
        situacion_fila = _situacion_facturacion(gasto)
        if situacion and situacion_fila != situacion:
            continue
        es_registro_absis = bool(getattr(gasto, "id_gasto_absis", None) or getattr(gasto, "datos_absis_json", None))
        filas.append({
            "id_expediente": exp.id_expediente if exp else getattr(gasto, "id_expediente_absis", None),
            "id_gasto_absis": getattr(gasto, "id_gasto_absis", None),
            "numero_factura_absis": getattr(gasto, "numero_factura_absis", None),
            "fecha_factura_absis": _fecha_iso(getattr(gasto, "fecha_factura_absis", None) or _fecha_original_absis(gasto, "FFACTURA")),
            "tipo_concepto_absis": getattr(gasto, "tipo_concepto_absis", None),
            "origen_registro": getattr(gasto, "origen_registro", None),
            "estado_consolidacion": getattr(gasto, "estado_consolidacion", None),
            "fecha_ultima_sincronizacion": _fecha_iso(getattr(gasto, "fecha_ultima_sincronizacion", None)),
            "tipo_cancelacion": tipo,
            "tipo_operacion": getattr(gasto, "tipo_operacion_absis", None) or (exp.tipo_operacion if exp else None),
            "concepto_gasto": getattr(gasto, "concepto_gasto_absis", None) if es_registro_absis else (getattr(gasto, "tipo", None) or getattr(gasto, "descripcion", None)),
            "situacion_facturacion": situacion_fila,
            "estado": getattr(gasto, "estado_absis", None),
            "estado_justificante": getattr(gasto, "estado_justificante", None),
            "fecha_adjuntar_factura": _fecha_iso(getattr(gasto, "fecha_adjuntar_factura", None)),
            "fecha_factura": _fecha_iso(getattr(gasto, "fecha_factura", None)),
            "imp_gasto": _decimal(getattr(gasto, "imp_gasto", None)) if es_registro_absis else _decimal(gasto.importe_total_bruto if gasto.importe_total_bruto is not None else gasto.importe),
            "imp_base_impuesto": _decimal(getattr(gasto, "imp_base_impuesto", None)) if es_registro_absis else _decimal(gasto.base_sujeta_impuestos),
            "imp_base_retencion": _decimal(getattr(gasto, "imp_base_retencion", None)) if es_registro_absis else _decimal(gasto.base_sujeta_retencion),
            "imp_impuesto": _decimal(getattr(gasto, "imp_impuesto", None)) if es_registro_absis else _decimal(gasto.importe_impuesto),
            "fecha_estado": _fecha_iso(getattr(gasto, "fecha_estado", None)),
            "fecha_peticion": _fecha_iso(getattr(gasto, "fecha_peticion", None)),
            "fecha_pago": _fecha_iso(getattr(gasto, "fecha_pago", None)),
            "estado_pago": gasto.estado_pago,
            "numero_factura": gasto.numero_factura,
            "datos_absis": (json.loads(gasto.datos_absis_json) if getattr(gasto, "datos_absis_json", None) else {}),
        })
    return filas

def _resumen(items):
    def suma(campo, seleccion=None):
        return round(sum(
            (x.get(campo) or 0) for x in items
            if seleccion is None or x.get("situacion_facturacion") == seleccion
        ), 2)
    def cuenta(situacion):
        return sum(1 for x in items if x.get("situacion_facturacion") == situacion)
    return {
        "registros": len(items),
        "total_gasto": suma("imp_gasto"),
        "base_impuesto": suma("imp_base_impuesto"),
        "base_retencion": suma("imp_base_retencion"),
        "total_impuesto": suma("imp_impuesto"),
        "con_provision": sum(1 for x in items if x["tipo_cancelacion"] == "con_provision"),
        "sin_provision": sum(1 for x in items if x["tipo_cancelacion"] == "sin_provision"),
        "otros": sum(1 for x in items if x["tipo_cancelacion"] == "otros"),
        "cobrado": cuenta("cobrado"),
        "pendiente_cobro": cuenta("pendiente_cobro"),
        "pendiente_factura": cuenta("pendiente_factura"),
        "incidencias": cuenta("incidencia"),
        "total_cobrado": suma("imp_gasto", "cobrado"),
        "total_pendiente_cobro": suma("imp_gasto", "pendiente_cobro"),
        "total_pendiente_factura": suma("imp_gasto", "pendiente_factura"),
        "total_incidencias": suma("imp_gasto", "incidencia"),
    }

@router.get("/facturacion")
def informe_facturacion(
    fecha_desde: Optional[date] = Query(None),
    fecha_hasta: Optional[date] = Query(None),
    tipo_operacion: Optional[str] = Query(None),
    estado: Optional[str] = Query(None),
    estado_justificante: Optional[str] = Query(None),
    situacion: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
    modalidad: str = Query("todas"),
    db: Session = Depends(get_db),
):
    if modalidad not in {"todas", "con_provision", "sin_provision", "otros"}:
        raise HTTPException(status_code=400, detail="Modalidad de cancelación no válida.")
    if situacion not in {None, "cobrado", "pendiente_cobro", "pendiente_factura", "incidencia"}:
        raise HTTPException(status_code=400, detail="Situación de facturación no válida.")
    items = _consultar(db, fecha_desde, fecha_hasta, tipo_operacion, estado, estado_justificante, situacion, q, modalidad)
    return {"resumen": _resumen(items), "items": items}

@router.get("/facturacion/exportar-excel")
def exportar_informe_facturacion(
    fecha_desde: Optional[date] = Query(None),
    fecha_hasta: Optional[date] = Query(None),
    tipo_operacion: Optional[str] = Query(None),
    estado: Optional[str] = Query(None),
    estado_justificante: Optional[str] = Query(None),
    situacion: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
    modalidad: str = Query("todas"),
    db: Session = Depends(get_db),
):
    if modalidad not in {"todas", "con_provision", "sin_provision", "otros"}:
        raise HTTPException(status_code=400, detail="Modalidad de cancelación no válida.")
    if situacion not in {None, "cobrado", "pendiente_cobro", "pendiente_factura", "incidencia"}:
        raise HTTPException(status_code=400, detail="Situación de facturación no válida.")
    items = _consultar(db, fecha_desde, fecha_hasta, tipo_operacion, estado, estado_justificante, situacion, q, modalidad)
    wb = Workbook()
    ws = wb.active
    ws.title = "Honorarios gestoría"
    ws.freeze_panes = "A2"
    columnas = [
        ("Nº expediente", "id_expediente"), ("IDGASTO", "id_gasto_absis"),
        ("Nº factura ABSIS", "numero_factura_absis"), ("Fecha factura ABSIS", "fecha_factura_absis"),
        ("TIPOCONCEPTO", "tipo_concepto_absis"), ("Origen", "origen_registro"),
        ("Estado consolidación", "estado_consolidacion"), ("Última sincronización", "fecha_ultima_sincronizacion"),
        ("Clasificación", "tipo_cancelacion"),
        ("TIPOOPERACION", "tipo_operacion"), ("CONCEPTOGASTO", "concepto_gasto"),
        ("Situación facturación", "situacion_facturacion"),
        ("ESTADO", "estado"), ("ESTADOJUSTIFICANTE", "estado_justificante"),
        ("FECHAADJUNTARFACTURA", "fecha_adjuntar_factura"), ("IMPGASTO (€)", "imp_gasto"),
        ("IMPBASEIMPUESTO (€)", "imp_base_impuesto"), ("IMPBASERETENCION (€)", "imp_base_retencion"),
        ("IMPIMPUESTO (€)", "imp_impuesto"), ("FESTADO", "fecha_estado"),
        ("FPETICION", "fecha_peticion"), ("FPAGO", "fecha_pago"), ("Estado pago ERP", "estado_pago"),
        ("Nº factura", "numero_factura"),
    ]
    # Añadir todas las columnas originales del Excel ABSIS, aunque no tengan
    # una columna normalizada propia en el modelo de datos.
    claves_originales = []
    for item in items:
        for clave in (item.get("datos_absis") or {}).keys():
            if clave not in claves_originales:
                claves_originales.append(clave)
    columnas.extend((f"ABSIS · {clave}", f"raw:{clave}") for clave in claves_originales)

    for idx, (label, _) in enumerate(columnas, 1):
        celda = ws.cell(row=1, column=idx, value=label)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1D4ED8")
        celda.alignment = Alignment(vertical="center")
    for rownum, item in enumerate(items, 2):
        for idx, (_, field) in enumerate(columnas, 1):
            if field.startswith("raw:"):
                valor = (item.get("datos_absis") or {}).get(field[4:])
            else:
                valor = item.get(field)
            ws.cell(row=rownum, column=idx, value=valor)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{max(1, len(items)+1)}"
    for idx, (label, field) in enumerate(columnas, 1):
        ws.column_dimensions[get_column_letter(idx)].width = min(max(len(label)+3, 14), 34)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="informe_facturacion_honorarios_gestoria.xlsx"'})
