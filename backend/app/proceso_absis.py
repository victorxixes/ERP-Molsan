from datetime import datetime
from io import BytesIO
from typing import Optional

import openpyxl
from openpyxl.styles import Font, PatternFill
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import Column, DateTime, Integer, String, case, func, or_
from sqlalchemy.orm import Session

from backend.app.database import Base, get_db
from backend.app.auth.dependencies import get_current_user
from backend.app.auth.permissions import require_perm_modulo

def _permiso_utilidades(accion: str):
    def comprobar(usuario=Depends(get_current_user)):
        require_perm_modulo(usuario, "utilidades", accion)
        return usuario
    return comprobar


router = APIRouter(prefix="/proceso-absis", tags=["Proceso ABSIS"])


class EnvioAbsis(Base):
    __tablename__ = "envios_absis"

    id = Column(Integer, primary_key=True, index=True)
    expediente = Column(String(120), nullable=False, index=True)
    fecha_envio = Column(DateTime, nullable=True, index=True)
    id_absis = Column(String(160), nullable=True, index=True)
    tipo = Column(String(200), nullable=False, default="General", index=True)
    fecha_creacion = Column(DateTime, nullable=False, server_default=func.now())
    fecha_incidencia = Column(DateTime, nullable=True)
    tipo_incidencia = Column(String(1000), nullable=True)
    estado = Column(String(40), nullable=False, default="Enviado", index=True)
    detalle = Column(String(4000), nullable=True)


class EnvioAbsisCrear(BaseModel):
    expediente: str = Field(..., min_length=1, max_length=120)
    fecha_envio: Optional[datetime] = None
    id_absis: Optional[str] = None
    tipo: str = Field(default="General", max_length=200)
    fecha_incidencia: Optional[datetime] = None
    tipo_incidencia: Optional[str] = None
    estado: str = "Enviado"
    detalle: Optional[str] = None


def _item(row):
    return {
        "id": row.id,
        "expediente": row.expediente or "",
        "fecha_envio": row.fecha_envio.isoformat() if row.fecha_envio else None,
        "id_absis": row.id_absis or "",
        "tipo": row.tipo or "General",
        "fecha_creacion": row.fecha_creacion.isoformat() if row.fecha_creacion else None,
        "fecha_incidencia": row.fecha_incidencia.isoformat() if row.fecha_incidencia else None,
        "tipo_incidencia": row.tipo_incidencia or "",
        "estado": row.estado or "Enviado",
        "detalle": row.detalle or "",
    }


def _consulta(db: Session, tipo: Optional[str], q: Optional[str]):
    query = db.query(EnvioAbsis)
    if tipo and tipo.strip():
        query = query.filter(EnvioAbsis.tipo == tipo.strip())
    if q and q.strip():
        term = f"%{q.strip()}%"
        query = query.filter(or_(
            EnvioAbsis.expediente.ilike(term),
            EnvioAbsis.id_absis.ilike(term),
            EnvioAbsis.tipo.ilike(term),
            EnvioAbsis.tipo_incidencia.ilike(term),
            EnvioAbsis.estado.ilike(term),
        ))
    return query


@router.get("/actividades")
def listar_actividades_absis(
    db: Session = Depends(get_db),
    _usuario=Depends(_permiso_utilidades("ver")),
    usuario=Depends(get_current_user),
):
    filas = db.query(
        EnvioAbsis.tipo,
        func.count(EnvioAbsis.id).label("enviados"),
        func.sum(case((EnvioAbsis.tipo_incidencia.isnot(None), 1), else_=0)).label("errores"),
    ).group_by(EnvioAbsis.tipo).order_by(EnvioAbsis.tipo.asc()).all()
    return [
        {"tipo": row.tipo or "General", "enviados": int(row.enviados or 0), "errores": int(row.errores or 0)}
        for row in filas
    ]


@router.get("/envios")
def listar_envios_absis(
    tipo: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    _usuario=Depends(_permiso_utilidades("ver")),
    usuario=Depends(get_current_user),
):
    rows = _consulta(db, tipo, q).order_by(EnvioAbsis.fecha_creacion.desc(), EnvioAbsis.id.desc()).all()
    return [_item(row) for row in rows]


@router.post("/envios", status_code=201)
def crear_envio_absis(
    datos: EnvioAbsisCrear,
    db: Session = Depends(get_db),
    _usuario=Depends(_permiso_utilidades("editar")),
    usuario=Depends(get_current_user),
):
    row = EnvioAbsis(
        expediente=datos.expediente.strip(),
        fecha_envio=datos.fecha_envio,
        id_absis=datos.id_absis,
        tipo=datos.tipo.strip() or "General",
        fecha_incidencia=datos.fecha_incidencia,
        tipo_incidencia=datos.tipo_incidencia,
        estado=datos.estado,
        detalle=datos.detalle,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return _item(row)


@router.delete("/envios/{envio_id}")
def eliminar_envio_absis(
    envio_id: int,
    db: Session = Depends(get_db),
    _usuario=Depends(_permiso_utilidades("editar")),
    usuario=Depends(get_current_user),
):
    row = db.query(EnvioAbsis).filter(EnvioAbsis.id == envio_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="No se ha encontrado el registro de envío.")
    db.delete(row)
    db.commit()
    return {"ok": True, "id": envio_id}


@router.post("/envios/{envio_id}/reintentar")
def reintentar_envio_absis(
    envio_id: int,
    db: Session = Depends(get_db),
    _usuario=Depends(_permiso_utilidades("editar")),
    usuario=Depends(get_current_user),
):
    row = db.query(EnvioAbsis).filter(EnvioAbsis.id == envio_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="No se ha encontrado el registro de envío.")
    # El proyecto aún no dispone de un cliente operativo para volver a transmitir a ABSIS.
    # Esta acción deja el envío en la cola lógica para reintento, sin simular un envío real.
    row.estado = "Pendiente de reintento"
    row.tipo_incidencia = row.tipo_incidencia or "Reintento solicitado manualmente"
    row.fecha_incidencia = datetime.utcnow()
    db.commit()
    db.refresh(row)
    return {"ok": True, "mensaje": "Envío marcado para reintento. La retransmisión requiere conectar el servicio ABSIS.", "envio": _item(row)}


@router.post("/envios/{envio_id}/pasado")
def marcar_envio_pasado_absis(
    envio_id: int,
    db: Session = Depends(get_db),
    _usuario=Depends(_permiso_utilidades("editar")),
    usuario=Depends(get_current_user),
):
    row = db.query(EnvioAbsis).filter(EnvioAbsis.id == envio_id).first()
    if not row:
        raise HTTPException(status_code=404, detail="No se ha encontrado el registro de envío.")
    row.estado = "Pasado"
    db.commit()
    db.refresh(row)
    return _item(row)


@router.get("/envios/exportar/excel")
def exportar_envios_absis(
    tipo: Optional[str] = Query(None),
    q: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    _usuario=Depends(_permiso_utilidades("ver")),
    usuario=Depends(get_current_user),
):
    rows = _consulta(db, tipo, q).order_by(EnvioAbsis.fecha_creacion.desc(), EnvioAbsis.id.desc()).all()
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Envíos ABSIS"
    hoja.append(["Expediente", "Fecha de envío", "ID", "IDAbsis", "Tipo", "Fecha de creación", "Fecha de incidencia", "Tipo de incidencia", "Estado", "Detalle"])
    for cell in hoja[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
    for row in rows:
        hoja.append([
            row.expediente, row.fecha_envio.isoformat(sep=" ") if row.fecha_envio else "",
            row.id, row.id_absis or "", row.tipo or "",
            row.fecha_creacion.isoformat(sep=" ") if row.fecha_creacion else "",
            row.fecha_incidencia.isoformat(sep=" ") if row.fecha_incidencia else "",
            row.tipo_incidencia or "", row.estado or "", row.detalle or "",
        ])
    for column in hoja.columns:
        letter = column[0].column_letter
        hoja.column_dimensions[letter].width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 14), 38)
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="envios_absis.xlsx"'},
    )
