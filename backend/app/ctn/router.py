import io
import os
import base64
import unicodedata
import hashlib
import secrets
from datetime import datetime


from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from cryptography.fernet import Fernet, InvalidToken
from sqlalchemy import Column, ForeignKey, Integer, Table, Float, and_, case, cast, func, literal, or_, select, text
from sqlalchemy.orm import aliased
from sqlalchemy.orm import Session
from openpyxl import Workbook

from backend.app.database import Base, get_db, engine
from backend.app.config import settings
from backend.app.auth.dependencies import get_current_user
from backend.app.auth.permissions import require_perm_modulo
from backend.app.ctn.models import Notaria
from backend.app.ctn.service import obtener_notaria
from backend.app.ctn.schemas import NotariaResponse, NotariaDatosBancariosUpdate
from backend.app.agenda.models import Cita
from backend.app.expedientes.models import Expediente
from backend.app.utils.distancia import distancia_km, MOLSAN_LAT, MOLSAN_LNG
from backend.app.ctn.geocode import geocode_todas_notarias
from backend.app.ctn.migracion import agregar_coordenadas
from backend.app.encargagest.models import EncargoExternoUsuario

router = APIRouter(
    prefix="/ctn",
    tags=["CTN"],
    # Todas las operaciones CTN requieren una sesión JWT válida.
    dependencies=[Depends(get_current_user)],
)


def _es_administrador_portales(usuario):
    rol = getattr(usuario, "rol", None)
    nombre = str(getattr(rol, "nombre", "") or "").strip().casefold()
    return nombre in {"administrador", "admin", "administrador del sistema"}


def _tiene_permiso_portal(usuario, portal, accion="ver"):
    # El administrador conserva la administración integral de los portales.
    if _es_administrador_portales(usuario):
        return True
    try:
        require_perm_modulo(usuario, portal, accion)
        return True
    except HTTPException:
        return False


def _exigir_permiso_portal(usuario, portal, accion="ver"):
    if not _tiene_permiso_portal(usuario, portal, accion):
        raise HTTPException(
            status_code=403,
            detail=f"No tienes permiso para '{accion}' en el portal '{portal}'. Configúralo en la ficha del empleado.",
        )


# ============================================================
# TABLA DE ASOCIACIONES RECÍPROCAS
# ============================================================
if "ctn_notarios_asociados" in Base.metadata.tables:
    notarios_asociados_table = Base.metadata.tables["ctn_notarios_asociados"]
else:
    notarios_asociados_table = Table(
        "ctn_notarios_asociados",
        Base.metadata,
        Column(
            "notario_id", Integer,
            ForeignKey("ctn_notarios.id", ondelete="CASCADE"),
            primary_key=True,
        ),
        Column(
            "asociado_id", Integer,
            ForeignKey("ctn_notarios.id", ondelete="CASCADE"),
            primary_key=True,
        ),
    )


@router.on_event("startup")
def crear_tabla_asociaciones_ctn():
    """Crea las asociaciones CTN y migra IBAN solo si faltan, serializando el DDL."""
    import time

    notarios_asociados_table.create(bind=engine, checkfirst=True)
    clave_bloqueo_migracion = 731208640224  # Compartida con las migraciones de main.py.
    conexion = None
    bloqueo_adquirido = False
    try:
        # No ejecutar ALTER TABLE en cada arranque: PostgreSQL necesita un bloqueo
        # exclusivo incluso cuando ADD COLUMN IF NOT EXISTS termina sin cambios.
        with engine.connect() as consulta:
            columnas = {
                fila[0]
                for fila in consulta.execute(text(
                    "SELECT column_name FROM information_schema.columns "
                    "WHERE table_schema = 'public' AND table_name = 'ctn_notarios'"
                )).fetchall()
            }
        pendientes = [nombre for nombre in ("iban_1", "iban_2") if nombre not in columnas]
        if not pendientes:
            print("[CTN DATOS BANCARIOS] Columnas IBAN disponibles; no se requiere DDL.", flush=True)
            return

        # Varias instancias pueden arrancar a la vez. Usamos el mismo advisory lock
        # global que main.py y volvemos a inspeccionar el esquema una vez adquirido.
        conexion = engine.connect()
        for intento in range(1, 61):
            with conexion.begin():
                bloqueo_adquirido = bool(
                    conexion.execute(
                        text("SELECT pg_try_advisory_lock(:clave)"),
                        {"clave": clave_bloqueo_migracion},
                    ).scalar()
                )
            if bloqueo_adquirido:
                break
            time.sleep(1)
        if not bloqueo_adquirido:
            print(
                "[CTN DATOS BANCARIOS] Se pospone la migración IBAN: "
                "otra inicialización mantiene el bloqueo global.",
                flush=True,
            )
            return

        with conexion.begin():
            conexion.execute(text("SET LOCAL lock_timeout = '3s'"))
            conexion.execute(text("SET LOCAL statement_timeout = '15s'"))
            columnas = {
                fila[0]
                for fila in conexion.execute(text(
                    "SELECT column_name FROM information_schema.columns "
                    "WHERE table_schema = 'public' AND table_name = 'ctn_notarios'"
                )).fetchall()
            }
            if "iban_1" not in columnas:
                conexion.execute(text(
                    "ALTER TABLE public.ctn_notarios ADD COLUMN iban_1 VARCHAR(34)"
                ))
            if "iban_2" not in columnas:
                conexion.execute(text(
                    "ALTER TABLE public.ctn_notarios ADD COLUMN iban_2 VARCHAR(34)"
                ))
        print("[CTN DATOS BANCARIOS] Columnas IBAN disponibles.", flush=True)
    except Exception as error_migracion:
        print(
            f"[CTN DATOS BANCARIOS] No se pudieron comprobar/crear las columnas IBAN: "
            f"{type(error_migracion).__name__}: {error_migracion}",
            flush=True,
        )
    finally:
        if conexion is not None:
            if bloqueo_adquirido:
                try:
                    with conexion.begin():
                        conexion.execute(
                            text("SELECT pg_advisory_unlock(:clave)"),
                            {"clave": clave_bloqueo_migracion},
                        )
                except Exception as error_unlock:
                    print(
                        f"[CTN DATOS BANCARIOS] No se pudo liberar el bloqueo: {error_unlock!r}.",
                        flush=True,
                    )
                    try:
                        conexion.invalidate()
                    except Exception:
                        pass
            conexion.close()


# ============================================================
# FUNCIONES AUXILIARES
# ============================================================
def _nombre_completo(notaria):
    partes = [getattr(notaria, "nombre", None) or "", getattr(notaria, "apellidos", None) or ""]
    return " ".join(parte.strip() for parte in partes if parte and parte.strip()).strip()


def _notaria_basica(notaria):
    data = NotariaResponse.from_orm(notaria).dict()
    data["nombre_completo"] = _nombre_completo(notaria)
    data["notargest_habilitado"] = bool(getattr(notaria, "notargest_habilitado", False))
    return data


def _asociaciones_de(db: Session, notaria_id: int):
    filas = (
        db.query(notarios_asociados_table.c.asociado_id)
        .filter(notarios_asociados_table.c.notario_id == notaria_id)
        .all()
    )
    ids = [fila[0] for fila in filas]
    if not ids:
        return []
    notarias = db.query(Notaria).filter(Notaria.id.in_(ids)).all()
    por_id = {notaria.id: notaria for notaria in notarias}
    return [
        {
            "id": por_id[ident].id,
            "nombre": por_id[ident].nombre,
            "apellidos": por_id[ident].apellidos,
            "nombre_completo": _nombre_completo(por_id[ident]),
            "nif": por_id[ident].nif,
        }
        for ident in ids
        if ident in por_id
    ]


def _notaria_con_asociaciones(db: Session, notaria):
    data = _notaria_basica(notaria)
    asociadas = _asociaciones_de(db, notaria.id)
    data["notarios_asociados"] = asociadas
    data["notario_asociado"] = ", ".join(
        item["nombre_completo"] or item["nif"] or str(item["id"])
        for item in asociadas
    )
    return data


def _condicion_videoconferencia():
    """Agrupa todos los valores históricos que representan VC.

    El importador antiguo guardaba VideoConferencia/Presencial, mientras
    que algunos Excels y el filtro de la UI utilizan SI/NO.
    """
    valor = func.lower(func.trim(func.coalesce(Notaria.vc, "")))
    return or_(
        valor.ilike("%videoconfer%"),
        valor.ilike("%video%confer%"),
        valor.in_(("si", "sí", "s", "yes", "true", "1", "vc", "video", "videoconference")),
    )


def _es_filtro_vc_positivo(valor: str) -> bool:
    normalizado = unicodedata.normalize("NFKD", valor.strip().lower())
    normalizado = "".join(c for c in normalizado if not unicodedata.combining(c))
    compacto = normalizado.replace(" ", "").replace("-", "").replace("_", "")
    return compacto in {"si", "s", "yes", "true", "1", "vc", "video", "videoconferencia", "videoconference"}


def _es_filtro_vc_negativo(valor: str) -> bool:
    normalizado = unicodedata.normalize("NFKD", valor.strip().lower())
    normalizado = "".join(c for c in normalizado if not unicodedata.combining(c))
    compacto = normalizado.replace(" ", "").replace("-", "").replace("_", "")
    return compacto in {"no", "n", "false", "0", "presencial"}


def _sin_acentos_sql(expresion):
    """Normaliza texto en PostgreSQL sin depender de la extensión unaccent."""
    return func.translate(
        func.lower(func.coalesce(expresion, "")),
        "áàäâãéèëêíìïîóòöôõúùüûñç",
        "aaaaaeeeeiiiiooooouuuunc",
    )


def _aplicar_filtros_notarias(
    db: Session,
    provincia=None,
    municipio=None,
    vc=None,
    apoderado=None,
    q=None,
    asociado=None,
):
    query = db.query(Notaria)

    if provincia and provincia.strip():
        query = query.filter(
            _sin_acentos_sql(Notaria.provincia).ilike(
                _sin_acentos_sql(f"%{provincia.strip()}%")
            )
        )

    if municipio and municipio.strip():
        query = query.filter(
            _sin_acentos_sql(Notaria.municipio).ilike(
                _sin_acentos_sql(f"%{municipio.strip()}%")
            )
        )

    if vc and vc.strip():
        filtro_vc = vc.strip()
        if _es_filtro_vc_positivo(filtro_vc):
            query = query.filter(_condicion_videoconferencia())
        elif _es_filtro_vc_negativo(filtro_vc):
            query = query.filter(~_condicion_videoconferencia())
        else:
            query = query.filter(
                func.lower(func.trim(func.coalesce(Notaria.vc, "")))
                == filtro_vc.lower()
            )

    if apoderado and apoderado.strip():
        termino_apoderado = f"%{apoderado.strip()}%"
        query = query.filter(
            _sin_acentos_sql(Notaria.apoderado).ilike(
                _sin_acentos_sql(termino_apoderado)
            )
        )

    if q and q.strip():
        # Buscar por palabras permite encontrar nombres compuestos aunque
        # nombre y apellidos estén guardados en columnas distintas.
        texto_busqueda = unicodedata.normalize("NFKD", q.strip().lower())
        texto_busqueda = "".join(
            caracter for caracter in texto_busqueda
            if not unicodedata.combining(caracter)
        )
        palabras = [
            palabra for palabra in texto_busqueda.split()
            if palabra not in {"notario", "notaria", "de", "del"}
        ]
        campos_busqueda = (
            _sin_acentos_sql(Notaria.nombre),
            _sin_acentos_sql(Notaria.apellidos),
            _sin_acentos_sql(Notaria.nif),
            _sin_acentos_sql(Notaria.codigo),
            _sin_acentos_sql(Notaria.municipio),
            _sin_acentos_sql(Notaria.provincia),
            _sin_acentos_sql(
                func.concat(
                    func.coalesce(Notaria.nombre, ""),
                    " ",
                    func.coalesce(Notaria.apellidos, ""),
                )
            ),
        )
        # Cada palabra puede aparecer en cualquier campo, y en cualquier orden.
        for palabra in palabras:
            termino = f"%{palabra}%"
            query = query.filter(
                or_(*[
                    campo.ilike(_sin_acentos_sql(termino))
                    for campo in campos_busqueda
                ])
            )

    if asociado and asociado.strip():
        termino_asociado = f"%{asociado.strip()}%"
        ids_coincidentes = (
            db.query(notarios_asociados_table.c.notario_id)
            .join(Notaria, Notaria.id == notarios_asociados_table.c.asociado_id)
            .filter(
                or_(
                    _sin_acentos_sql(Notaria.nombre).ilike(_sin_acentos_sql(termino_asociado)),
                    _sin_acentos_sql(Notaria.apellidos).ilike(_sin_acentos_sql(termino_asociado)),
                    _sin_acentos_sql(Notaria.nif).ilike(_sin_acentos_sql(termino_asociado)),
                )
            )
            .distinct()
            .subquery()
        )
        query = query.filter(Notaria.id.in_(db.query(ids_coincidentes.c.notario_id)))

    return query


def _clasificar_tipo_firma(tipo_firma):
    valor = unicodedata.normalize("NFKD", str(tipo_firma or ""))
    valor = "".join(c for c in valor if not unicodedata.combining(c)).strip().lower()
    compacto = valor.replace(" ", "").replace("-", "_").replace("_", "")
    if not valor:
        return "SIN_CLASIFICAR"
    if compacto in {"vc", "video", "videoconferencia", "videoconference"}:
        return "VC"
    if "videoconferencia" in compacto or "videoconference" in compacto:
        return "VC"
    if "presencial" in compacto or compacto in {"p", "fisica"}:
        return "PRESENCIAL"
    return "SIN_CLASIFICAR"


def _valor_fecha(valor):
    return valor.isoformat() if valor else None


def _distancia_de_notaria(notaria):
    try:
        lat = getattr(notaria, "lat", None)
        lng = getattr(notaria, "lng", None)
        if lat is None or lng is None:
            return None
        return distancia_km(MOLSAN_LAT, MOLSAN_LNG, float(lat), float(lng))
    except (TypeError, ValueError):
        return None


# ============================================================
# GEOCODIFICACIÓN Y MIGRACIÓN
# ============================================================
@router.post("/geocode/notarias")
def geocode_notarias(db: Session = Depends(get_db)):
    return geocode_todas_notarias(db)


@router.post("/migracion/agregar-coordenadas")
def migrar_coordenadas(db: Session = Depends(get_db)):
    return agregar_coordenadas(db)


# ============================================================
# LISTADO Y EXPORTACIÓN
# ============================================================
@router.get("/notarias")
def listar_notarias(
    provincia: str | None = None,
    municipio: str | None = None,
    vc: str | None = None,
    apoderado: str | None = None,
    q: str | None = None,
    asociado: str | None = Query(None, description="Filtra por nombre, apellidos o NIF del notario asociado"),
    page: int = Query(1, ge=1),
    page_size: int = Query(15, ge=1, le=200),
    sort_by: str = Query("apellidos", max_length=40),
    sort_dir: str = Query("asc", regex="^(asc|desc)$"),
    db: Session = Depends(get_db),
):
    query = _aplicar_filtros_notarias(db, provincia, municipio, vc, apoderado, q, asociado)
    total = query.count()

    # Ordenación en base de datos antes de paginar: nunca se ordena solo la
    # página visible. Los campos calculados también tienen expresión SQL.
    asociada = aliased(Notaria)
    asociados_nombre = (
        select(func.coalesce(func.string_agg(
            func.trim(func.concat(func.coalesce(asociada.nombre, ""), " ", func.coalesce(asociada.apellidos, "")),
            ", "
        ), ""))
        .select_from(notarios_asociados_table.join(asociada, asociada.id == notarios_asociados_table.c.asociado_id))
        .where(notarios_asociados_table.c.notario_id == Notaria.id)
        .scalar_subquery()
    )
    acceso_encargagest = (
        select(EncargoExternoUsuario.habilitado)
        .where(EncargoExternoUsuario.notaria_id == Notaria.id)
        .scalar_subquery()
    )

    lat_valida = Notaria.lat.op("~")(r"^-?[0-9]+(\.[0-9]+)?$")
    lng_valida = Notaria.lng.op("~")(r"^-?[0-9]+(\.[0-9]+)?$")
    latitud = case((lat_valida, cast(Notaria.lat, Float)), else_=None)
    longitud = case((lng_valida, cast(Notaria.lng, Float)), else_=None)
    phi1 = func.radians(literal(MOLSAN_LAT))
    phi2 = func.radians(latitud)
    delta_phi = func.radians(latitud - literal(MOLSAN_LAT))
    delta_lambda = func.radians(longitud - literal(MOLSAN_LNG))
    haversine = (
        func.pow(func.sin(delta_phi / 2), 2)
        + func.cos(phi1) * func.cos(phi2) * func.pow(func.sin(delta_lambda / 2), 2)
    )
    distancia_sql = case(
        (and_(lat_valida, lng_valida), 6371.0 * 2 * func.asin(func.sqrt(func.least(1.0, haversine)))),
        else_=None,
    )

    nombre_completo_sql = func.trim(func.concat(func.coalesce(Notaria.nombre, ""), " ", func.coalesce(Notaria.apellidos, "")))
    expresiones_orden = {
        "id": Notaria.id,
        "nombre_completo": nombre_completo_sql,
        "nif": Notaria.nif,
        "municipio": Notaria.municipio,
        "provincia": Notaria.provincia,
        "vc": Notaria.vc,
        "apoderado": Notaria.apoderado,
        "notario_asociado": asociados_nombre,
        "notargest_habilitado": Notaria.notargest_habilitado,
        "encargagest_habilitado": func.coalesce(acceso_encargagest, False),
        "distancia_km": distancia_sql,
        "apellidos": Notaria.apellidos,
    }
    expresion = expresiones_orden.get(sort_by, Notaria.apellidos)
    orden = expresion.desc().nullslast() if sort_dir.lower() == "desc" else expresion.asc().nullslast()
    filas = (
        query.order_by(orden, Notaria.id.asc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )

    accesos_encarga = {
        int(notaria_id): bool(habilitado)
        for notaria_id, habilitado in db.query(
            EncargoExternoUsuario.notaria_id, EncargoExternoUsuario.habilitado
        ).filter(EncargoExternoUsuario.notaria_id.in_([fila.id for fila in filas])).all()
    } if filas else {}
    items = []
    for notaria in filas:
        data = _notaria_con_asociaciones(db, notaria)
        data["distancia_km"] = _distancia_de_notaria(notaria)
        data["encargagest_habilitado"] = accesos_encarga.get(int(notaria.id), False)
        items.append(data)
    return {"items": items, "total": total, "page": page, "page_size": page_size, "sort_by": sort_by if sort_by in expresiones_orden else "apellidos", "sort_dir": sort_dir.lower()}


@router.get("/exportar-excel")
def exportar_notarias_excel(
    provincia: str | None = None,
    municipio: str | None = None,
    vc: str | None = None,
    apoderado: str | None = None,
    q: str | None = None,
    asociado: str | None = None,
    db: Session = Depends(get_db),
):
    query = _aplicar_filtros_notarias(db, provincia, municipio, vc, apoderado, q, asociado)
    filas = query.order_by(Notaria.apellidos.asc().nullslast(), Notaria.nombre.asc().nullslast()).all()
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Notarías"
    hoja.append([
        "Código", "Nombre", "Apellidos", "NIF", "Teléfono", "Provincia", "Municipio",
        "CP", "Dirección", "Videoconferencia", "Apoderado", "Departamento cancelaciones",
        "Departamento copias", "Otros departamentos", "Notarios asociados", "Observaciones", "Distancia km",
    ])
    for notaria in filas:
        nombres_asociados = "; ".join(
            item["nombre_completo"] or item["nif"] or str(item["id"])
            for item in _asociaciones_de(db, notaria.id)
        )
        hoja.append([
            getattr(notaria, "codigo", None), getattr(notaria, "nombre", None),
            getattr(notaria, "apellidos", None), getattr(notaria, "nif", None),
            getattr(notaria, "telefono", None), getattr(notaria, "provincia", None),
            getattr(notaria, "municipio", None), getattr(notaria, "cp", None),
            getattr(notaria, "direccion", None), getattr(notaria, "vc", None),
            getattr(notaria, "apoderado", None), getattr(notaria, "departamento_cancelaciones", None),
            getattr(notaria, "departamento_copias", None), getattr(notaria, "otros_departamentos", None),
            nombres_asociados, getattr(notaria, "observacion", None), _distancia_de_notaria(notaria),
        ])
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    for columna in hoja.columns:
        letra = columna[0].column_letter
        longitud_maxima = max((len(str(celda.value)) for celda in columna if celda.value is not None), default=10)
        hoja.column_dimensions[letra].width = min(max(longitud_maxima + 2, 12), 45)
    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="notarias.xlsx"'},
    )


# ============================================================
# FICHA Y ASOCIACIONES RECÍPROCAS
# ============================================================
@router.get("/notarias/{notaria_id}")
def detalle_notaria(notaria_id: int, db: Session = Depends(get_db)):
    notaria = obtener_notaria(db, notaria_id)
    if notaria is None:
        raise HTTPException(status_code=404, detail="Notario no encontrado")
    return _notaria_con_asociaciones(db, notaria)


@router.put("/notarias/{notaria_id}/datos-bancarios")
def actualizar_datos_bancarios_notaria(
    notaria_id: int,
    payload: NotariaDatosBancariosUpdate,
    db: Session = Depends(get_db),
):
    """Guarda hasta dos IBAN para reutilizarlos al registrar pagos del notario."""
    notaria = db.query(Notaria).filter(Notaria.id == notaria_id).first()
    if notaria is None:
        raise HTTPException(status_code=404, detail="Notario no encontrado")
    for campo, valor in payload.dict().items():
        normalizado = str(valor or "").replace(" ", "").replace("\u00a0", "").strip().upper()
        setattr(notaria, campo, normalizado or None)
    try:
        db.commit()
        db.refresh(notaria)
    except Exception:
        db.rollback()
        raise HTTPException(status_code=500, detail="No se pudieron guardar los datos bancarios del notario.")
    return _notaria_con_asociaciones(db, notaria)


@router.get("/notarias/{notaria_id}/asociadas")
def listar_notarias_asociadas(notaria_id: int, db: Session = Depends(get_db)):
    notaria = db.query(Notaria).filter(Notaria.id == notaria_id).first()
    if notaria is None:
        raise HTTPException(status_code=404, detail="Notario no encontrado")
    return {"notaria_id": notaria_id, "items": _asociaciones_de(db, notaria_id)}


@router.post("/notarias/{notaria_id}/asociadas/{asociada_id}")
def asociar_notarios(notaria_id: int, asociada_id: int, db: Session = Depends(get_db)):
    if notaria_id == asociada_id:
        raise HTTPException(status_code=400, detail="Un notario no puede asociarse consigo mismo")
    notaria = db.query(Notaria).filter(Notaria.id == notaria_id).first()
    asociada = db.query(Notaria).filter(Notaria.id == asociada_id).first()
    if notaria is None or asociada is None:
        raise HTTPException(status_code=404, detail="No se ha encontrado uno de los notarios")
    try:
        for origen, destino in ((notaria_id, asociada_id), (asociada_id, notaria_id)):
            existe = (
                db.query(notarios_asociados_table.c.notario_id)
                .filter(and_(
                    notarios_asociados_table.c.notario_id == origen,
                    notarios_asociados_table.c.asociado_id == destino,
                ))
                .first()
            )
            if not existe:
                db.execute(notarios_asociados_table.insert().values(notario_id=origen, asociado_id=destino))
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {
        "ok": True,
        "notario_id": notaria_id,
        "asociado": {"id": asociada.id, "nombre_completo": _nombre_completo(asociada), "nif": asociada.nif},
        "items": _asociaciones_de(db, notaria_id),
    }


@router.delete("/notarias/{notaria_id}/asociadas/{asociada_id}")
def desasociar_notarios(notaria_id: int, asociada_id: int, db: Session = Depends(get_db)):
    notaria = db.query(Notaria).filter(Notaria.id == notaria_id).first()
    asociada = db.query(Notaria).filter(Notaria.id == asociada_id).first()
    if notaria is None or asociada is None:
        raise HTTPException(status_code=404, detail="No se ha encontrado uno de los notarios")
    try:
        db.execute(notarios_asociados_table.delete().where(or_(
            and_(notarios_asociados_table.c.notario_id == notaria_id, notarios_asociados_table.c.asociado_id == asociada_id),
            and_(notarios_asociados_table.c.notario_id == asociada_id, notarios_asociados_table.c.asociado_id == notaria_id),
        )))
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"ok": True, "notario_id": notaria_id, "asociada_id": asociada_id}


# ============================================================
# RESUMEN LEGADO DE CITAS DE AGENDA (NO SON FIRMAS REALIZADAS)
# ============================================================
@router.get("/notarias/{notaria_id}/firmas")
def obtener_resumen_citas_notario(notaria_id: int, db: Session = Depends(get_db)):
    notaria = db.query(Notaria).filter(Notaria.id == notaria_id).first()
    if notaria is None:
        raise HTTPException(status_code=404, detail="Notario no encontrado")
    total = db.query(func.count(Cita.id)).filter(Cita.notario_id == notaria_id).scalar() or 0
    total_vc = db.query(func.count(Cita.id)).filter(Cita.notario_id == notaria_id, Cita.tipo_cita == "VC").scalar() or 0
    total_presencial = db.query(func.count(Cita.id)).filter(Cita.notario_id == notaria_id, Cita.tipo_cita == "P").scalar() or 0
    return {
        "notaria_id": notaria_id,
        "total_firmas": total,
        "total_vc": total_vc,
        "total_presencial": total_presencial,
        "tipo_dato": "citas_agenda_no_firmas_realizadas",
    }


# ============================================================
# HISTÓRICO REAL DE FIRMAS: EXPEDIENTES CON FECHA_FIRMA
# ============================================================
@router.get("/notarias/{notaria_id}/firmas/historico")
def historico_firmas_notario(
    notaria_id: int,
    modalidad: str | None = Query(None, description="VC o PRESENCIAL"),
    q: str | None = None,
    db: Session = Depends(get_db),
):
    notaria = db.query(Notaria).filter(Notaria.id == notaria_id).first()
    if notaria is None:
        raise HTTPException(status_code=404, detail="Notario no encontrado")

    condiciones = [Expediente.fecha_firma.isnot(None)]
    # CORRECCIÓN: si hay NIF, se filtra por NIF; solo se usa el nombre como fallback si no hay NIF.
    nif_notaria = (notaria.nif or "").strip()
    if nif_notaria:
        condiciones.append(
            func.lower(func.trim(func.coalesce(Expediente.nif_notario, ""))) == nif_notaria.lower()
        )
    else:
        nombre = _nombre_completo(notaria)
        if not nombre:
            return {
                "notaria_id": notaria_id, "notario": "", "nif_notario": notaria.nif,
                "items": [], "total": 0, "total_vc": 0, "total_presencial": 0,
                "total_sin_clasificar": 0,
            }
        condiciones.append(or_(
            func.lower(func.trim(func.coalesce(Expediente.nombre_notario, ""))) == nombre.lower(),
            func.lower(func.trim(func.coalesce(Expediente.notario, ""))) == nombre.lower(),
        ))

    expedientes = (
        db.query(Expediente).filter(*condiciones)
        .order_by(Expediente.fecha_firma.desc(), Expediente.id.desc()).all()
    )
    modalidad_filtro = (modalidad or "").strip().upper()
    if modalidad_filtro not in {"", "VC", "VIDEOCONFERENCIA", "P", "PRESENCIAL"}:
        raise HTTPException(status_code=400, detail="Modalidad debe ser VC o PRESENCIAL")

    items = []
    for expediente in expedientes:
        tipo = _clasificar_tipo_firma(getattr(expediente, "tipo_firma", None))
        if modalidad_filtro in {"VC", "VIDEOCONFERENCIA"} and tipo != "VC":
            continue
        if modalidad_filtro in {"P", "PRESENCIAL"} and tipo != "PRESENCIAL":
            continue
        campos_busqueda = [
            getattr(expediente, "id_expediente", None), getattr(expediente, "nombre_titular", None),
            getattr(expediente, "nif_titular", None), getattr(expediente, "nombre_solicitante", None),
            getattr(expediente, "nif_solicitante", None), getattr(expediente, "protocolo", None),
            getattr(expediente, "oficina", None),
        ]
        if q and q.strip():
            texto_busqueda = " ".join(str(v or "") for v in campos_busqueda).lower()
            if q.strip().lower() not in texto_busqueda:
                continue
        items.append({
            "id": expediente.id,
            "id_expediente": getattr(expediente, "id_expediente", None),
            "fecha_firma": _valor_fecha(getattr(expediente, "fecha_firma", None)),
            "tipo_firma": getattr(expediente, "tipo_firma", None),
            "modalidad": tipo,
            "protocolo": getattr(expediente, "protocolo", None),
            "nombre_titular": getattr(expediente, "nombre_titular", None),
            "nif_titular": getattr(expediente, "nif_titular", None),
            "nombre_solicitante": getattr(expediente, "nombre_solicitante", None),
            "nif_solicitante": getattr(expediente, "nif_solicitante", None),
            "oficina": getattr(expediente, "oficina", None),
            "capital": getattr(expediente, "capital", None),
            "importe": getattr(expediente, "importe", None),
            "estado_expediente": getattr(expediente, "estado_expediente", None),
            "observaciones": getattr(expediente, "observaciones", None),
        })

    return {
        "notaria_id": notaria_id,
        "notario": _nombre_completo(notaria),
        "nif_notario": notaria.nif,
        "items": items,
        "total": len(items),
        "total_vc": sum(1 for item in items if item["modalidad"] == "VC"),
        "total_presencial": sum(1 for item in items if item["modalidad"] == "PRESENCIAL"),
        "total_sin_clasificar": sum(1 for item in items if item["modalidad"] == "SIN_CLASIFICAR"),
    }


# ACCESO INDIVIDUAL A NOTARGEST
@router.on_event("startup")
def migrar_columna_acceso_notargest():
    """Añade el indicador de acceso sin dejar que un bloqueo auxiliar tumbe la API."""
    import time as _time

    try:
        with engine.connect() as conexion:
            existe = conexion.execute(
                text(
                    "SELECT 1 FROM information_schema.columns "
                    "WHERE table_schema = 'public' "
                    "AND table_name = 'ctn_notarios' "
                    "AND column_name = 'notargest_habilitado'"
                )
            ).first()

        if existe:
            print("[NOTARGEST MIGRACION] notargest_habilitado ya existe; se omite ALTER.", flush=True)
            return

        for intento in range(1, 4):
            try:
                with engine.begin() as conexion:
                    conexion.execute(text("SET LOCAL lock_timeout = '5s'"))
                    conexion.execute(text("SET LOCAL statement_timeout = '15s'"))
                    conexion.execute(
                        text(
                            "ALTER TABLE public.ctn_notarios "
                            "ADD COLUMN notargest_habilitado BOOLEAN NOT NULL DEFAULT FALSE"
                        )
                    )
                print("[NOTARGEST MIGRACION] Columna notargest_habilitado creada.", flush=True)
                return
            except Exception as error:
                print(
                    f"[NOTARGEST MIGRACION] Intento {intento}/3 fallido: {error!r}",
                    flush=True,
                )
                if intento < 3:
                    _time.sleep(2)

        print(
            "[NOTARGEST MIGRACION] AVISO: no se pudo añadir notargest_habilitado; "
            "el arranque continuará. Reintentar la migración cuando desaparezca el bloqueo.",
            flush=True,
        )
    except Exception as error:
        print(
            f"[NOTARGEST MIGRACION] AVISO: no se pudo comprobar la columna; "
            f"el arranque continuará y habrá que revisar la migración: {error!r}",
            flush=True,
        )


@router.put("/notarias/{notaria_id}/notargest")
def cambiar_acceso_notargest(
    notaria_id: int,
    habilitado: bool = Query(...),
    db: Session = Depends(get_db),
    usuario=Depends(get_current_user),
):
    """Habilita o revoca el acceso al portal con permiso explícito de edición."""
    _exigir_permiso_portal(usuario, "notargestweb", "editar")
    notaria = db.query(Notaria).filter(Notaria.id == notaria_id).first()
    if not notaria:
        raise HTTPException(status_code=404, detail="Notaría no encontrada")
    notaria.notargest_habilitado = bool(habilitado)
    _registrar_evento_portal(db, notaria_id, "notargest", "CAMBIO_ACCESO",
        detalle="Acceso habilitado" if habilitado else "Acceso deshabilitado")
    db.commit()
    db.refresh(notaria)
    return {
        "ok": True,
        "notaria_id": notaria.id,
        "habilitado": bool(notaria.notargest_habilitado),
        "mensaje": "Acceso a NotarGest habilitado" if habilitado else "Acceso a NotarGest deshabilitado",
    }


@router.on_event("startup")
def crear_tabla_historial_accesos_portales():
    """Prepara el historial sin bloquear el arranque si PostgreSQL tiene un lock."""
    def ejecutar_migracion(etiqueta, sql, parametros=None):
        try:
            with engine.begin() as conexion:
                # Evita que una migración auxiliar deje Render esperando hasta el timeout.
                conexion.execute(text("SET LOCAL lock_timeout = '3s'"))
                conexion.execute(text("SET LOCAL statement_timeout = '8s'"))
                conexion.execute(text(sql), parametros or {})
            print(f"[CTN PORTALES] {etiqueta}: OK.", flush=True)
            return True
        except Exception as exc:
            # Estas tablas auxiliares no deben impedir que arranque toda la API.
            print(f"[CTN PORTALES] AVISO: {etiqueta} no se pudo completar; se reintentará en otro arranque. Error: {exc!r}", flush=True)
            return False

    print("[CTN PORTALES] Preparando historial y cifrado de contraseñas...", flush=True)
    tabla_ok = ejecutar_migracion(
        "crear tabla de historial",
        """
        CREATE TABLE IF NOT EXISTS ctn_portal_accesos_historial (
            id SERIAL PRIMARY KEY, notaria_id INTEGER NOT NULL,
            portal VARCHAR(30) NOT NULL, evento VARCHAR(40) NOT NULL,
            resultado VARCHAR(20) NOT NULL DEFAULT 'ok', usuario VARCHAR(150),
            detalle VARCHAR(500), fecha TIMESTAMP NOT NULL DEFAULT CURRENT_TIMESTAMP
        )
        """,
    )
    if tabla_ok:
        ejecutar_migracion(
            "crear índice del historial",
            """
            CREATE INDEX IF NOT EXISTS ix_ctn_portal_historial_notaria_fecha
            ON ctn_portal_accesos_historial (notaria_id, fecha DESC)
            """,
        )

    # La columna es independiente del historial: se migra en otra transacción
    # para que un bloqueo de una tabla no retenga el resto de la inicialización.
    # No ejecutar ALTER TABLE en cada arranque si la columna ya existe:
    # incluso ADD COLUMN IF NOT EXISTS necesita un bloqueo de tabla en PostgreSQL.
    try:
        with engine.connect() as conexion:
            columna_password_enc = conexion.execute(
                text(
                    "SELECT 1 FROM information_schema.columns "
                    "WHERE table_schema = 'public' "
                    "AND table_name = 'ctn_notarios' "
                    "AND column_name = 'notargest_password_enc'"
                )
            ).first()
    except Exception as exc:
        print(
            "[CTN PORTALES] AVISO: no se pudo comprobar "
            "ctn_notarios.notargest_password_enc; se reintentará en otro arranque.",
            flush=True,
        )
        columna_password_enc = None

    if columna_password_enc:
        print(
            "[CTN PORTALES] ctn_notarios.notargest_password_enc ya existe; se omite ALTER.",
            flush=True,
        )
    else:
        ejecutar_migracion(
            "añadir columna ctn_notarios.notargest_password_enc",
            """
            ALTER TABLE public.ctn_notarios
            ADD COLUMN IF NOT EXISTS notargest_password_enc TEXT
            """,
        )
    print("[CTN PORTALES] Revisión de esquema finalizada; la API puede continuar.", flush=True)

def _registrar_evento_portal(db, notaria_id, portal, evento, resultado="ok", usuario=None, detalle=None):
    db.execute(text("""
        INSERT INTO ctn_portal_accesos_historial
            (notaria_id, portal, evento, resultado, usuario, detalle, fecha)
        VALUES (:notaria_id, :portal, :evento, :resultado, :usuario, :detalle, CURRENT_TIMESTAMP)
    """), {"notaria_id": notaria_id, "portal": portal, "evento": evento,
           "resultado": resultado, "usuario": usuario, "detalle": detalle})



def _fernet_passwords():
    secreto = str(getattr(settings, "JWT_SECRET", "") or os.getenv("JWT_SECRET", "")).strip()
    if not secreto:
        raise HTTPException(status_code=500, detail="No está configurada la clave de cifrado de contraseñas.")
    clave = base64.urlsafe_b64encode(hashlib.sha256(("ERP-Molsan-portales-v1:" + secreto).encode("utf-8")).digest())
    return Fernet(clave)


def _cifrar_password_portal(password):
    return _fernet_passwords().encrypt(password.encode("utf-8")).decode("ascii")


def _descifrar_password_portal(password_cifrada):
    try:
        return _fernet_passwords().decrypt(password_cifrada.encode("ascii")).decode("utf-8")
    except (InvalidToken, ValueError, TypeError, UnicodeDecodeError):
        raise HTTPException(status_code=500, detail="No se ha podido descifrar la contraseña almacenada. Comprueba la clave JWT_SECRET.")

def _hash_password_portal(password):
    salt = secrets.token_hex(16)
    rounds = 310000
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), bytes.fromhex(salt), rounds).hex()
    return "pbkdf2_sha256$%s$%s$%s" % (rounds, salt, digest)


def _normalizar_usuario_portal(valor):
    return "".join(str(valor or "").upper().split()).replace("-", "").replace(".", "")


@router.get("/accesos-portales/resumen")
def resumen_accesos_portales_notarias(
    db: Session = Depends(get_db),
    usuario=Depends(get_current_user),
):
    puede_notargest = _tiene_permiso_portal(usuario, "notargestweb", "ver")
    puede_encargagest = _tiene_permiso_portal(usuario, "encargagestweb", "ver")
    if not (puede_notargest or puede_encargagest):
        raise HTTPException(status_code=403, detail="No tienes permiso para consultar los accesos notariales.")

    total_notargest = None
    total_encargagest = None
    if puede_notargest:
        total_notargest = db.query(Notaria).filter(
            Notaria.notargest_habilitado.is_(True)
        ).count()
    if puede_encargagest:
        total_encargagest = db.query(func.count(func.distinct(EncargoExternoUsuario.notaria_id))).filter(
            EncargoExternoUsuario.habilitado.is_(True)
        ).scalar() or 0

    return {
        "notargest_habilitados": total_notargest,
        "encargagest_habilitados": total_encargagest,
        "permisos": {
            "notargestweb": {"ver": puede_notargest},
            "encargagestweb": {"ver": puede_encargagest},
        },
    }


@router.get("/notarias/{notaria_id}/accesos-portales")
def obtener_accesos_portales_notaria(
    notaria_id: int,
    db: Session = Depends(get_db),
    usuario=Depends(get_current_user),
):
    puede_notargest = _tiene_permiso_portal(usuario, "notargestweb", "ver")
    puede_encargagest = _tiene_permiso_portal(usuario, "encargagestweb", "ver")
    if not (puede_notargest or puede_encargagest):
        raise HTTPException(status_code=403, detail="No tienes permiso para consultar los accesos a NotarGestWeb o EncargaGestWeb.")
    notaria = db.query(Notaria).filter(Notaria.id == notaria_id).first()
    if not notaria:
        raise HTTPException(status_code=404, detail="Notaría no encontrada")
    acceso_eg = db.query(EncargoExternoUsuario).filter(EncargoExternoUsuario.notaria_id == notaria_id).first()
    historial = db.execute(text("""
        SELECT portal, evento, resultado, usuario, detalle, fecha
        FROM ctn_portal_accesos_historial WHERE notaria_id = :notaria_id
        ORDER BY fecha DESC, id DESC LIMIT 100
    """), {"notaria_id": notaria_id}).mappings().all()
    return {
        "notaria_id": notaria_id,
        "notargest": ({"usuario": _normalizar_usuario_portal(notaria.nif),
            "habilitado": bool(getattr(notaria, "notargest_habilitado", False)),
            "configurado": bool(getattr(notaria, "notargest_password_hash", None)),
            "cambiar_password": bool(getattr(notaria, "notargest_cambiar_password", True))}
            if puede_notargest else None),
        "encargagest": ({"usuario": acceso_eg.usuario if acceso_eg else _normalizar_usuario_portal(notaria.nif),
            "habilitado": bool(acceso_eg.habilitado) if acceso_eg else False,
            "configurado": acceso_eg is not None}
            if puede_encargagest else None),
        "permisos": {"notargestweb": {"ver": puede_notargest}, "encargagestweb": {"ver": puede_encargagest}},
        "historial": [{"portal": row["portal"], "evento": row["evento"],
            "resultado": row["resultado"], "usuario": row["usuario"], "detalle": row["detalle"],
            "fecha": row["fecha"].isoformat() if row["fecha"] else None} for row in historial
            if (row["portal"] == "notargest" and puede_notargest) or (row["portal"] == "encargagest" and puede_encargagest)],
    }



@router.get("/notarias/{notaria_id}/accesos-portales/{portal}/password")
def revelar_password_portal_notaria(
    notaria_id: int,
    portal: str,
    db: Session = Depends(get_db),
    usuario=Depends(get_current_user),
):
    portal = portal.strip().lower()
    if not _es_administrador_portales(usuario):
        raise HTTPException(status_code=403, detail="Solo un administrador puede revelar las contraseñas de los portales.")
    clave_permiso = {"notargest": "notargestweb", "encargagest": "encargagestweb"}.get(portal)
    if not clave_permiso:
        raise HTTPException(status_code=400, detail="Portal no válido")
    _exigir_permiso_portal(usuario, clave_permiso, "ver")
    notaria = db.query(Notaria).filter(Notaria.id == notaria_id).first()
    if not notaria:
        raise HTTPException(status_code=404, detail="Notaría no encontrada")
    if portal == "notargest":
        password_cifrada = getattr(notaria, "notargest_password_enc", None)
        usuario_login = _normalizar_usuario_portal(notaria.nif)
    else:
        acceso = db.query(EncargoExternoUsuario).filter(EncargoExternoUsuario.notaria_id == notaria_id).first()
        password_cifrada = getattr(acceso, "password_enc", None) if acceso else None
        usuario_login = acceso.usuario if acceso else _normalizar_usuario_portal(notaria.nif)
    if not password_cifrada:
        raise HTTPException(status_code=409, detail="Esta contraseña es anterior al almacenamiento cifrado y no se puede recuperar. Restablécela una vez para poder mostrarla con el ojo.")
    password = _descifrar_password_portal(password_cifrada)
    _registrar_evento_portal(db, notaria_id, portal, "REVELAR_PASSWORD",
        detalle="Un administrador reveló la contraseña")
    db.commit()
    return {"ok": True, "portal": portal, "usuario": usuario_login, "password": password}


@router.post("/notarias/{notaria_id}/accesos-portales/{portal}/habilitado")
def cambiar_habilitado_portal_notaria(
    notaria_id: int,
    portal: str,
    datos: dict,
    db: Session = Depends(get_db),
    usuario=Depends(get_current_user),
):
    portal = portal.strip().lower()
    clave_permiso = {"notargest": "notargestweb", "encargagest": "encargagestweb"}.get(portal)
    if not clave_permiso:
        raise HTTPException(status_code=400, detail="Portal no válido")
    _exigir_permiso_portal(usuario, clave_permiso, "editar")
    notaria = db.query(Notaria).filter(Notaria.id == notaria_id).first()
    if not notaria:
        raise HTTPException(status_code=404, detail="Notaría no encontrada")
    habilitado = bool(datos.get("habilitado"))
    if portal == "notargest":
        notaria.notargest_habilitado = habilitado
    elif portal == "encargagest":
        acceso = db.query(EncargoExternoUsuario).filter(EncargoExternoUsuario.notaria_id == notaria_id).first()
        if not acceso:
            raise HTTPException(status_code=409, detail="Primero genera las credenciales de EncargaGest.")
        acceso.habilitado = habilitado
        acceso.actualizado_en = datetime.utcnow()
    else:
        raise HTTPException(status_code=400, detail="Portal no válido")
    _registrar_evento_portal(db, notaria_id, portal, "CAMBIO_ACCESO",
        detalle="Acceso habilitado" if habilitado else "Acceso deshabilitado")
    db.commit()
    return {"ok": True, "portal": portal, "habilitado": habilitado}


@router.post("/notarias/{notaria_id}/accesos-portales/{portal}/reset-password")
def reset_password_portal_notaria(
    notaria_id: int,
    portal: str,
    db: Session = Depends(get_db),
    usuario=Depends(get_current_user),
):
    portal = portal.strip().lower()
    clave_permiso = {"notargest": "notargestweb", "encargagest": "encargagestweb"}.get(portal)
    if not clave_permiso:
        raise HTTPException(status_code=400, detail="Portal no válido")
    notaria = db.query(Notaria).filter(Notaria.id == notaria_id).first()
    if not notaria:
        raise HTTPException(status_code=404, detail="Notaría no encontrada")
    if portal == "notargest":
        accion_permiso = "editar" if getattr(notaria, "notargest_password_hash", None) else "crear"
    else:
        acceso_existente = db.query(EncargoExternoUsuario).filter(
            EncargoExternoUsuario.notaria_id == notaria_id
        ).first()
        accion_permiso = "editar" if acceso_existente else "crear"
    _exigir_permiso_portal(usuario, clave_permiso, accion_permiso)
    password_nueva = _normalizar_usuario_portal(notaria.nif)
    usuario_login = _normalizar_usuario_portal(notaria.nif)
    if not usuario_login:
        raise HTTPException(status_code=422, detail="La notaría debe tener NIF registrado en CTN.")
    if portal == "notargest":
        notaria.notargest_password_hash = _hash_password_portal(password_nueva)
        notaria.notargest_password_enc = _cifrar_password_portal(password_nueva)
        notaria.notargest_habilitado = True
        notaria.notargest_cambiar_password = True
        notaria.notargest_password_provisional = True
    elif portal == "encargagest":
        acceso = db.query(EncargoExternoUsuario).filter(EncargoExternoUsuario.notaria_id == notaria_id).first()
        duplicado = db.query(EncargoExternoUsuario).filter(
            func.upper(EncargoExternoUsuario.usuario) == usuario_login,
            EncargoExternoUsuario.notaria_id != notaria_id).first()
        if duplicado:
            raise HTTPException(status_code=409, detail="El usuario basado en el NIF ya está asignado a otra notaría.")
        if acceso is None:
            acceso = EncargoExternoUsuario(notaria_id=notaria_id, usuario=usuario_login,
                password_hash=_hash_password_portal(password_nueva),
                password_enc=_cifrar_password_portal(password_nueva), password_provisional=True, habilitado=True)
            db.add(acceso)
        else:
            acceso.usuario = usuario_login
            acceso.password_hash = _hash_password_portal(password_nueva)
            acceso.password_enc = _cifrar_password_portal(password_nueva)
            acceso.password_provisional = True
            acceso.habilitado = True
            acceso.actualizado_en = datetime.utcnow()
    else:
        raise HTTPException(status_code=400, detail="Portal no válido")
    _registrar_evento_portal(db, notaria_id, portal, "RESET_PASSWORD",
        detalle="Contraseña provisional restablecida al NIF registrado en CTN; la notaría puede cambiarla desde su perfil")
    db.commit()
    return {"ok": True, "portal": portal, "usuario": usuario_login,
        "password_nueva": password_nueva, "habilitado": True,
        "mensaje": "Contraseña provisional restablecida al NIF del notario. Puede cambiarla desde Mi perfil."}
