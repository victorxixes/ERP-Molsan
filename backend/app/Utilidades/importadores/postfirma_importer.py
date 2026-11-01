"""Carga pura del histórico ABSIS Post Firma.

Cada fichero y cada fila se conserva de forma independiente. No se crean ni
actualizan gastos ERP, no se vinculan expedientes y no se deduplican movimientos.
"""
import hashlib
import json
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session


def _limpiar(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, (datetime, date)):
        return value
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def _json_value(value):
    value = _limpiar(value)
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if hasattr(value, "item"):
        try:
            return value.item()
        except Exception:
            pass
    return value if isinstance(value, (str, int, float, bool)) else str(value)


def _normalizar_columna(value):
    return "".join(
        caracter for caracter in str(value or "").strip().upper()
        if caracter.isalnum()
    )


def _texto(value):
    value = _limpiar(value)
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def _decimal(value):
    value = _limpiar(value)
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip().replace(" ", "")
        if not value:
            return None
        if "," in value:
            value = value.replace(".", "").replace(",", ".")
    try:
        result = Decimal(str(value))
        return result if result.is_finite() else None
    except (InvalidOperation, ValueError, TypeError):
        return None


def _fecha(value):
    value = _limpiar(value)
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            numero = float(value)
            if 20000 <= numero <= 80000:
                return (datetime(1899, 12, 30) + timedelta(days=numero)).date()
        except (TypeError, ValueError, OverflowError):
            return None
    try:
        parsed = pd.to_datetime(value, dayfirst=True, errors="coerce")
        return None if pd.isna(parsed) else parsed.date()
    except Exception:
        return None


def _crear_almacen(db: Session):
    """Crea tablas nuevas y aisladas; no altera tablas de gastos o expedientes."""
    db.execute(text("""
        CREATE TABLE IF NOT EXISTS public.absis_postfirma_cargas (
            id BIGSERIAL PRIMARY KEY,
            nombre_archivo TEXT NOT NULL,
            sha256 VARCHAR(64) NOT NULL,
            fecha_carga TIMESTAMP WITHOUT TIME ZONE NOT NULL DEFAULT CURRENT_TIMESTAMP,
            filas_excel INTEGER NOT NULL DEFAULT 0,
            filas_guardadas INTEGER NOT NULL DEFAULT 0,
            columnas_json TEXT NOT NULL DEFAULT '[]'
        )
    """))
    db.execute(text("""
        CREATE TABLE IF NOT EXISTS public.absis_postfirma_registros (
            id BIGSERIAL PRIMARY KEY,
            carga_id BIGINT NOT NULL REFERENCES public.absis_postfirma_cargas(id) ON DELETE RESTRICT,
            numero_fila INTEGER NOT NULL,
            datos_json TEXT NOT NULL,
            id_expediente_absis TEXT NULL,
            id_gasto_absis TEXT NULL,
            tipo_concepto TEXT NULL,
            concepto_gasto TEXT NULL,
            estado_absis TEXT NULL,
            importe NUMERIC(20, 4) NULL,
            fecha_pago DATE NULL
        )
    """))
    db.execute(text("CREATE INDEX IF NOT EXISTS ix_absis_postfirma_registros_carga ON public.absis_postfirma_registros(carga_id, numero_fila)"))
    db.execute(text("CREATE INDEX IF NOT EXISTS ix_absis_postfirma_registros_idgasto ON public.absis_postfirma_registros(id_gasto_absis)"))
    db.execute(text("CREATE INDEX IF NOT EXISTS ix_absis_postfirma_registros_idexpediente ON public.absis_postfirma_registros(id_expediente_absis)"))
    db.commit()


def importar_postfirma(db: Session, contenido_excel: bytes, nombre_archivo: str = ""):
    """Guarda cada fila con contenido de una hoja Excel como dato histórico bruto.

    Reimportar el mismo fichero genera otra carga independiente. No se eliminan
    duplicados: pueden representar movimientos legítimos o periodos distintos.
    """
    try:
        df = pd.read_excel(BytesIO(contenido_excel), sheet_name=0, header=0, engine="openpyxl", dtype=object)
        libro = load_workbook(BytesIO(contenido_excel), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"No se pudo leer el Excel Post Firma: {exc}")

    hoja = libro.worksheets[0]
    columnas = [str(c) if c is not None else "" for c in df.columns]
    if "TIPOCONCEPTO" not in [str(c).strip().upper() for c in columnas]:
        columnas.append("TIPOCONCEPTO")
    sha256 = hashlib.sha256(contenido_excel).hexdigest()
    guardadas = 0
    filas_vacias = 0
    try:
        _crear_almacen(db)
        carga = db.execute(text("""
            INSERT INTO public.absis_postfirma_cargas
                (nombre_archivo, sha256, filas_excel, filas_guardadas, columnas_json)
            VALUES (:nombre, :sha, :filas, 0, :columnas)
            RETURNING id
        """), {
            "nombre": nombre_archivo or "historico_postfirma.xlsx",
            "sha": sha256,
            "filas": int(len(df)),
            "columnas": json.dumps(columnas, ensure_ascii=False),
        }).first()
        carga_id = int(carga[0])

        # La carga y todas sus filas forman una única transacción: si falla un
        # lote, el rollback elimina también la cabecera y los lotes anteriores.
        # Iteración directa de las celdas para conservar valores originales de Excel.
        filas_originales = hoja.iter_rows(min_row=2, max_col=max(len(columnas), 1), values_only=False)
        lote = []
        for indice, row in df.iterrows():
            numero_fila = int(indice) + 2
            celdas = next(filas_originales, ())
            datos = {}
            for posicion, nombre in enumerate(df.columns):
                etiqueta = str(nombre) if nombre is not None else ""
                if posicion >= len(celdas):
                    datos[etiqueta] = None
                else:
                    datos[etiqueta] = _json_value(celdas[posicion].value)
            valores_normalizados = {
                _normalizar_columna(k): v for k, v in datos.items()
            }
            datos["TIPOCONCEPTO"] = valores_normalizados.get("TIPOCONCEPTO")
            if not any(v is not None and v != "" for v in datos.values()):
                filas_vacias += 1
                continue

            # Estos campos son índices de consulta auxiliares, no reglas de importación.
            normalizados = {_normalizar_columna(k): v for k, v in datos.items()}
            importe = _decimal(normalizados.get("IMPGASTO"))
            fecha_pago = _fecha(normalizados.get("FPAGO"))
            lote.append({
                "carga_id": carga_id,
                "numero_fila": numero_fila,
                "datos_json": json.dumps(datos, ensure_ascii=False, default=str),
                "id_expediente": _texto(normalizados.get("IDEXPEDIENTE")),
                "id_gasto": _texto(normalizados.get("IDGASTO")),
                "tipo_concepto": _texto(normalizados.get("TIPOCONCEPTO")),
                "concepto_gasto": _texto(normalizados.get("CONCEPTOGASTO")),
                "estado": _texto(normalizados.get("ESTADO")),
                "importe": importe,
                "fecha_pago": fecha_pago,
            })
            if len(lote) >= 500:
                db.execute(text("""
                    INSERT INTO public.absis_postfirma_registros
                    (carga_id, numero_fila, datos_json, id_expediente_absis, id_gasto_absis,
                     tipo_concepto, concepto_gasto, estado_absis, importe, fecha_pago)
                    VALUES (:carga_id, :numero_fila, :datos_json, :id_expediente, :id_gasto,
                            :tipo_concepto, :concepto_gasto, :estado, :importe, :fecha_pago)
                """), lote)
                guardadas += len(lote)
                lote = []

        if lote:
            db.execute(text("""
                INSERT INTO public.absis_postfirma_registros
                (carga_id, numero_fila, datos_json, id_expediente_absis, id_gasto_absis,
                 tipo_concepto, concepto_gasto, estado_absis, importe, fecha_pago)
                VALUES (:carga_id, :numero_fila, :datos_json, :id_expediente, :id_gasto,
                        :tipo_concepto, :concepto_gasto, :estado, :importe, :fecha_pago)
            """), lote)
            guardadas += len(lote)
        db.execute(text("""
            UPDATE public.absis_postfirma_cargas
            SET filas_guardadas = :guardadas WHERE id = :id
        """), {"guardadas": guardadas, "id": carga_id})
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        libro.close()

    return {
        "ok": True,
        "carga_id": carga_id,
        "archivo": nombre_archivo,
        "sha256": sha256,
        "filas_excel": int(len(df)),
        "filas_guardadas": guardadas,
        "filas_vacias": filas_vacias,
        "columnas_detectadas": columnas,
        "nota": "Carga histórica aislada. No se han creado ni actualizado gastos ERP ni vinculado expedientes. No se han deduplicado filas.",
    }


def listar_cargas_postfirma(db: Session):
    _crear_almacen(db)
    filas = db.execute(text("""
        SELECT id, nombre_archivo, sha256, fecha_carga, filas_excel, filas_guardadas, columnas_json
        FROM public.absis_postfirma_cargas ORDER BY id DESC
    """)).mappings().all()
    return [{
        "id": int(f["id"]),
        "nombre_archivo": f["nombre_archivo"],
        "sha256": f["sha256"],
        "fecha_carga": f["fecha_carga"].isoformat() if f["fecha_carga"] else None,
        "filas_excel": f["filas_excel"],
        "filas_guardadas": f["filas_guardadas"],
        "columnas": json.loads(f["columnas_json"] or "[]"),
    } for f in filas]


def listar_registros_postfirma(
    db: Session,
    limit: int = 200,
    offset: int = 0,
    carga_id: int | None = None,
    sort_by: str = "carga_id",
    sort_dir: str = "desc",
):
    _crear_almacen(db)
    if carga_id is not None:
        esquema = db.execute(text(
            "SELECT columnas_json FROM public.absis_postfirma_cargas WHERE id = :id"
        ), {"id": int(carga_id)}).first()
        columnas_informe = json.loads(esquema[0] or "[]") if esquema else []
    else:
        esquemas = db.execute(text(
            "SELECT columnas_json FROM public.absis_postfirma_cargas ORDER BY id DESC"
        )).all()
        columnas_informe = list(dict.fromkeys(
            columna
            for esquema in esquemas
            for columna in json.loads(esquema[0] or "[]")
        ))

    # El campo dinámico se pasa como parámetro JSONB, nunca como identificador
    # SQL. Solo se permite ordenar por columnas originales del Excel o metadatos.
    campos_sistema = {"carga_id", "archivo", "numero_fila", "fecha_carga"}
    campo = sort_by if sort_by in campos_sistema or sort_by in columnas_informe else "carga_id"
    direccion = "DESC" if str(sort_dir).lower() == "desc" else "ASC"
    where = "WHERE r.carga_id = :carga_id" if carga_id is not None else ""
    params = {"limit": max(1, min(int(limit), 500)), "offset": max(0, int(offset))}
    if carga_id is not None:
        params["carga_id"] = int(carga_id)
    if campo == "carga_id":
        order_sql = f"r.carga_id {direccion}"
    elif campo == "archivo":
        order_sql = f"LOWER(c.nombre_archivo) {direccion}"
    elif campo == "numero_fila":
        order_sql = f"r.numero_fila {direccion}"
    elif campo == "fecha_carga":
        order_sql = f"c.fecha_carga {direccion} NULLS LAST"
    else:
        params["sort_column"] = campo
        valor_sql = "(r.datos_json::jsonb ->> :sort_column)"
        # Orden numérico para valores Excel numéricos; texto natural para el resto.
        order_sql = (
            f"CASE WHEN {valor_sql} ~ '^[+-]?[0-9]+([.][0-9]+)?$' "
            f"THEN ({valor_sql})::numeric END {direccion} NULLS LAST, "
            f"LOWER(COALESCE({valor_sql}, '')) {direccion}"
        )
    params_count = {key: value for key, value in params.items() if key != "sort_column"}
    total = db.execute(text(
        f"SELECT COUNT(*) FROM public.absis_postfirma_registros r {where}"
    ), params_count).scalar() or 0
    filas = db.execute(text(f"""
        SELECT r.id, r.carga_id, r.numero_fila, r.datos_json, c.nombre_archivo, c.fecha_carga
        FROM public.absis_postfirma_registros r
        JOIN public.absis_postfirma_cargas c ON c.id = r.carga_id
        {where}
        ORDER BY {order_sql}, r.carga_id DESC, r.numero_fila ASC
        LIMIT :limit OFFSET :offset
    """), params).mappings().all()
    return {
        "total": int(total),
        "limit": params["limit"],
        "offset": params["offset"],
        "columnas": columnas_informe,
        "sort_by": campo,
        "sort_dir": direccion.lower(),
        "items": [{
            "id": int(f["id"]),
            "carga_id": int(f["carga_id"]),
            "numero_fila": int(f["numero_fila"]),
            "archivo": f["nombre_archivo"],
            "fecha_carga": f["fecha_carga"].isoformat() if f["fecha_carga"] else None,
            "datos": json.loads(f["datos_json"] or "{}"),
        } for f in filas],
    }


def validar_postfirma(db: Session, contenido_excel: bytes, fecha_desde=None, fecha_hasta=None):
    """Previsualización sin escritura; no impone fechas ni encabezados obligatorios."""
    try:
        df = pd.read_excel(BytesIO(contenido_excel), sheet_name=0, header=0, engine="openpyxl", dtype=object)
    except Exception as exc:
        raise ValueError(f"No se pudo leer el Excel Post Firma: {exc}")
    columnas = [str(c) if c is not None else "" for c in df.columns]
    if "TIPOCONCEPTO" not in [str(c).strip().upper() for c in columnas]:
        columnas.append("TIPOCONCEPTO")
    filas_con_datos = 0
    filas_vacias = 0
    importes = Decimal("0")
    filas_con_importe = 0
    estados = {}
    for _, row in df.iterrows():
        valores = [_limpiar(v) for v in row.tolist()]
        if not any(v is not None and v != "" for v in valores):
            filas_vacias += 1
            continue
        filas_con_datos += 1
        datos = {_normalizar_columna(k): v for k, v in zip(columnas, valores)}
        importe = _decimal(datos.get("IMPGASTO"))
        if importe is not None:
            importes += importe
            filas_con_importe += 1
        estado = _texto(datos.get("ESTADO")) or "(vacío)"
        estados[estado] = estados.get(estado, 0) + 1
    return {
        "validacion_sin_escritura": True,
        "columnas_detectadas": columnas,
        "filas_excel": int(len(df)),
        "filas_con_datos": filas_con_datos,
        "filas_vacias": filas_vacias,
        "filas_con_importe_valido": filas_con_importe,
        "suma_imp_gasto": str(importes.quantize(Decimal("0.01"))),
        "por_estado_absis": estados,
        "nota": "Validación informativa; no filtra por fecha, estado, concepto ni IDGASTO.",
    }
