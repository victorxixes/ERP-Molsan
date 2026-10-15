from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from backend.app.ctn.models import Notaria


HEADER_MAP = {
    "Código": "codigo",
    "Nombre": "nombre",
    "Apellidos": "apellidos",
    "NIF": "nif",
    "Teléfono": "telefono",
    "Departamento cancelaciones": "departamento_cancelaciones",
    "Departamento copias": "departamento_copias",
    "Otros departamentos": "otros_departamentos",
    "CP": "cp",
    "Provincia": "provincia",
    "Municipio": "municipio",
    "Dirección": "direccion",
    "VC": "vc",
    "Apoderado": "apoderado",
    "Apoderado S": "apoderado_s",
    "Observación": "observacion",
}


def limpiar_texto(valor):
    if valor is None:
        return ""
    resultado = str(valor).strip()
    return "" if resultado.lower() == "none" else resultado


def normalizar_vc(valor) -> str:
    """Guarda solo los dos valores canónicos que usa el catálogo CTN."""
    texto = limpiar_texto(valor).casefold().replace("í", "i")
    compacto = texto.replace(" ", "").replace("-", "").replace("_", "")
    if compacto in {"si", "s", "vc", "videoconferencia", "videoconference", "true", "1"}:
        return "VideoConferencia"
    return "Presencial"


def importar_ctn_desde_excel(db: Session, contenido: bytes) -> int:
    """Importa CTN de forma transaccional.

    El archivo se lee y valida íntegramente antes de tocar datos existentes.
    Si falla cualquier inserción, se revierte también el borrado, preservando el
    catálogo previo. No se hace commit intermedio.
    """
    if not contenido:
        raise ValueError("El fichero Excel está vacío.")

    try:
        wb = load_workbook(BytesIO(contenido), read_only=True, data_only=True)
        ws = wb.active
        filas = ws.iter_rows(values_only=True)
        headers = None
        datos_importar = []

        for fila in filas:
            if fila and any(limpiar_texto(celda) == "Código" for celda in fila if celda is not None):
                headers = [limpiar_texto(celda) for celda in fila]
                break
        if headers is None:
            raise ValueError("No se encontraron cabeceras válidas: falta la columna 'Código'.")

        cabeceras_requeridas = {"Código", "Nombre", "Apellidos"}
        faltan = cabeceras_requeridas.difference(set(headers))
        if faltan:
            raise ValueError("Faltan columnas obligatorias: " + ", ".join(sorted(faltan)))

        for fila in filas:
            if not fila or all(celda is None for celda in fila):
                continue
            datos = {}
            for idx, header in enumerate(headers):
                campo = HEADER_MAP.get(header)
                if campo:
                    valor = fila[idx] if idx < len(fila) else None
                    datos[campo] = limpiar_texto(valor)
            # No insertar filas sin código: evita registros huérfanos al importar.
            if not datos.get("codigo"):
                continue
            datos["vc"] = normalizar_vc(datos.get("vc"))
            for campo in ("apoderado", "apoderado_s", "observacion"):
                datos[campo] = limpiar_texto(datos.get(campo))
            datos_importar.append(datos)
        wb.close()

        if not datos_importar:
            raise ValueError("El Excel no contiene notarías con código válido; no se han modificado datos.")

        # Una sola transacción: el rollback recupera los registros anteriores.
        db.query(Notaria).delete(synchronize_session=False)
        db.add_all([Notaria(**datos) for datos in datos_importar])
        db.commit()
        return len(datos_importar)

    except Exception:
        db.rollback()
        raise
