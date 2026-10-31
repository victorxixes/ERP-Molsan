from io import BytesIO
from datetime import date, datetime

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, or_, cast, String
from sqlalchemy.orm import Session

from backend.app.database import get_db
from backend.app.auth.dependencies import get_current_user
from backend.app.auth.permissions import require_perm_modulo
from backend.app.expediente_ppad.models import ExpedientePPAD
from backend.app.expedientes.models import Expediente, Cliente
from backend.app.empleados.models import Empleado


def _permiso_utilidades(usuario=Depends(get_current_user)):
    require_perm_modulo(usuario, "utilidades", "ver")
    return usuario


router = APIRouter(
    prefix="/utilidades/ppad",
    tags=["Utilidades - PPAD"],
    dependencies=[Depends(get_current_user)],
)


def _fecha(valor):
    return valor.isoformat() if valor else None


def _consultar(db: Session, estado: str = "todos", q: str = ""):
    consulta = (
        db.query(ExpedientePPAD, Expediente, Empleado, Cliente)
        .join(Expediente, Expediente.id == ExpedientePPAD.expediente_id)
        .outerjoin(Empleado, Empleado.id == Expediente.gestor_asignado_id)
        .outerjoin(Cliente, Cliente.id == Expediente.cliente_id)
    )

    if estado == "pendiente":
        consulta = consulta.filter(ExpedientePPAD.compensado.is_(False))
    elif estado == "compensado":
        consulta = consulta.filter(ExpedientePPAD.compensado.is_(True))

    termino = (q or "").strip()
    if termino:
        like = f"%{termino}%"
        consulta = consulta.filter(
            or_(
                cast(ExpedientePPAD.id, String).ilike(like),
                ExpedientePPAD.concepto.ilike(like),
                Expediente.id_expediente.ilike(like),
                func.coalesce(Expediente.gestor, "").ilike(like),
                func.coalesce(Expediente.nombre_solicitante, "").ilike(like),
                func.coalesce(Expediente.nombre_titular, "").ilike(like),
                func.coalesce(Cliente.nombre_completo, "").ilike(like),
                func.concat(
                    func.coalesce(Empleado.nombre, ""),
                    " ",
                    func.coalesce(Empleado.apellidos, ""),
                ).ilike(like),
            )
        )

    return consulta.order_by(ExpedientePPAD.fecha_creacion.desc(), ExpedientePPAD.id.desc()).all()


def _fila(item, expediente, gestor, cliente):
    nombre_gestor = " ".join(
        parte for parte in (getattr(gestor, "nombre", None), getattr(gestor, "apellidos", None)) if parte
    ).strip()
    solicitante = (
        getattr(expediente, "nombre_solicitante", None)
        or getattr(expediente, "nombre_titular", None)
        or getattr(cliente, "nombre_completo", None)
        or ""
    )
    return {
        "id": item.id,
        "expediente_id": expediente.id_expediente,
        "fecha_firma": _fecha(getattr(expediente, "fecha_firma", None)),
        "documento": f"PPAD-{item.id}",
        "fecha_alta": _fecha(item.fecha_creacion),
        "fecha_compensacion": _fecha(item.fecha_compensacion),
        "gestor": nombre_gestor or getattr(expediente, "gestor", None) or "—",
        "concepto": item.concepto or "",
        "solicitante": solicitante,
        "importe": float(item.importe or 0),
        "imputable_a": item.imputable_a or "Molsan",
        "compensado": bool(item.compensado),
    }


@router.get("", summary="Listado y resumen de PPAD pendientes y compensados")
def listar_ppad_utilidades(
    db: Session = Depends(get_db),
    q: str = Query("", max_length=200),
    _usuario=Depends(_permiso_utilidades),
):
    filas = _consultar(db, "todos", q)
    items = [_fila(*fila) for fila in filas]

    pendientes = [item for item in items if not item["compensado"]]
    compensados = [item for item in items if item["compensado"]]

    return {
        "items": items,
        "resumen": {
            "pendientes_total": len(pendientes),
            "pendientes_importe": round(sum(item["importe"] for item in pendientes), 2),
            "compensados_total": len(compensados),
            "compensados_importe": round(sum(item["importe"] for item in compensados), 2),
            "total": len(items),
        },
    }


@router.get("/exportar/excel", summary="Exportar PPAD a Excel")
def exportar_ppad_excel(
    estado: str = Query("pendiente", regex="^(pendiente|compensado|todos)$"),
    q: str = Query("", max_length=200),
    db: Session = Depends(get_db),
    _usuario=Depends(_permiso_utilidades),
):
    filas = _consultar(db, estado, q)
    items = [_fila(*fila) for fila in filas]
    libro = Workbook()
    hoja = libro.active
    hoja.title = "PPAD"
    hoja.append([
        "Expediente", "Documento", "Fecha alta", "Fecha compensación",
        "Gestor", "Concepto", "Solicitante", "Importe (€)", "Imputable a", "Estado",
    ])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1E40AF")

    for item in items:
        hoja.append([
            item["expediente_id"],
            item["documento"],
            item["fecha_alta"] or "",
            item["fecha_compensacion"] or "",
            item["gestor"],
            item["concepto"],
            item["solicitante"],
            item["importe"],
            item["imputable_a"],
            "Compensado" if item["compensado"] else "Pendiente de compensar",
        ])

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    for columna in hoja.columns:
        letra = columna[0].column_letter
        maximo = max((len(str(celda.value or "")) for celda in columna), default=10)
        hoja.column_dimensions[letra].width = min(max(maximo + 2, 12), 42)

    buffer = BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    nombre = f"ppad_{estado}_{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )
